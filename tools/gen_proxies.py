"""Generate ``src/aioopenpyxl/_generated.py`` and the ``*.pyi`` module shims.

Forwarders
    For every public member of openpyxl's Workbook / Worksheet /
    ReadOnlyWorksheet / WriteOnlyWorksheet / Chartsheet (as described by the
    types-openpyxl stubs) a typed, synchronous forwarder is emitted so the
    async wrappers expose the complete openpyxl surface with IDE-visible
    signatures.  Members that need special handling (blocking I/O, wrapping of
    returned sheets, dunders) are listed in ``EXCLUDE`` and written by hand in
    the wrapper classes.  Members that exist at runtime but are missing from
    the stubs are added through ``EXTRA_MEMBERS``.

Module shims
    ``aioopenpyxl._compat`` registers every ``openpyxl.*`` sub-module as
    ``aioopenpyxl.*`` in ``sys.modules`` at import time.  Type checkers cannot
    see that, so for every public openpyxl module that has a stub a
    ``.pyi``-only shim (``from openpyxl.<mod> import *``) is written next to the
    package.  No ``.py`` is placed there: at runtime the ``sys.modules`` alias
    wins, so the shim is invisible to the interpreter.

Self-checks (the script exits non-zero on the first, warns on the others)
    * signature drift: stub parameter names / defaults of every non-overloaded
      forwarder are compared with ``inspect.signature`` of the runtime method;
    * runtime members (``dir(instance)``) that neither the stubs, ``EXCLUDE``
      nor ``EXTRA_MEMBERS`` know about;
    * stub defaults written as ``...`` that cannot be resolved at runtime.

Run:  uv run python tools/gen_proxies.py
"""

from __future__ import annotations

import ast
import importlib
import importlib.metadata
import inspect
import io
import pkgutil
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
PKG_DIR = ROOT / "src" / "aioopenpyxl"
OUT = PKG_DIR / "_generated.py"


def stub_dir() -> Path:
    import openpyxl

    pkg = Path(openpyxl.__file__).parent
    stubs = pkg.parent / "openpyxl-stubs"
    if not stubs.is_dir():
        sys.exit("types-openpyxl is not installed (uv sync --group dev)")
    return stubs


def warn(msg: str) -> None:
    print(f"warning: {msg}", file=sys.stderr)


# (stub file, class name, runtime import path)
SPECS: dict[str, list[tuple[str, str, str]]] = {
    "WorkbookProxy": [
        ("workbook/workbook.pyi", "Workbook", "openpyxl.workbook.workbook:Workbook"),
    ],
    "WorksheetProxy": [
        ("workbook/child.pyi", "_WorkbookChild", "openpyxl.workbook.child:_WorkbookChild"),
        ("worksheet/worksheet.pyi", "Worksheet", "openpyxl.worksheet.worksheet:Worksheet"),
        (
            "worksheet/_read_only.pyi",
            "ReadOnlyWorksheet",
            "openpyxl.worksheet._read_only:ReadOnlyWorksheet",
        ),
        (
            "worksheet/_write_only.pyi",
            "WriteOnlyWorksheet",
            "openpyxl.worksheet._write_only:WriteOnlyWorksheet",
        ),
    ],
    "ChartsheetProxy": [
        ("workbook/child.pyi", "_WorkbookChild", "openpyxl.workbook.child:_WorkbookChild"),
        ("chartsheet/chartsheet.pyi", "Chartsheet", "openpyxl.chartsheet.chartsheet:Chartsheet"),
    ],
}

# Hand-written in the Async* classes.
EXCLUDE: dict[str, set[str]] = {
    "WorkbookProxy": {
        "active",
        "worksheets",
        "chartsheets",
        "sheetnames",
        "create_sheet",
        "create_chartsheet",
        "copy_worksheet",
        "remove",
        "remove_sheet",
        "index",
        "get_index",
        "move_sheet",
        "get_sheet_by_name",
        "get_sheet_names",
        "save",
        "close",
    },
    "WorksheetProxy": {
        "parent",
        "cell",  # read-only guard added by hand
        "iter_rows",
        "iter_cols",
        "rows",
        "columns",
        "values",
        "append",
        "calculate_dimension",
        "close",
    },
    "ChartsheetProxy": {"parent"},
}

# Instance attributes that exist at runtime but are missing from the stubs:
# name -> (annotation, TYPE_CHECKING import statement or None).  Emitted as
# read/write properties after the stub-derived members.
EXTRA_MEMBERS: dict[str, dict[str, tuple[str, str | None]]] = {
    "WorkbookProxy": {
        "custom_doc_props": (
            "CustomPropertyList[Any]",
            "from openpyxl.packaging.custom import CustomPropertyList",
        ),
    },
    "WorksheetProxy": {},
    "ChartsheetProxy": {
        # Inherited from openpyxl.descriptors.serialisable.Serialisable, whose
        # body the Chartsheet stub does not repeat.
        "idx_base": ("int", None),
        "namespace": ("str | None", None),
    },
}

# Overload sets copied verbatim from the stubs that pyright flags as overlapping
# (reportOverlappingOverload); the first overload gets a targeted ignore.
OVERLAPPING_OVERLOADS: dict[str, set[str]] = {
    "WorksheetProxy": {"cell"},
}

# Names imported at runtime by the generated module; dropped from the
# TYPE_CHECKING block so they are not imported twice.
RUNTIME_IMPORTS: dict[str, set[str]] = {
    "typing": {"TYPE_CHECKING", "Any", "overload"},
    "typing_extensions": {"deprecated"},
}

NEVER_TYPES = {"Never", "NoReturn", "typing.Never", "typing.NoReturn", "typing_extensions.Never"}

RAW_TYPES = {
    "WorkbookProxy": "Workbook",
    "WorksheetProxy": "Worksheet | ReadOnlyWorksheet | WriteOnlyWorksheet",
    "ChartsheetProxy": "Chartsheet",
}


@dataclass
class Prop:
    name: str
    getter: str  # return annotation
    setter: str | None = None  # value annotation
    deprecated: str | None = None  # decorator source, e.g. deprecated('...')


@dataclass
class Method:
    name: str
    defs: list[ast.FunctionDef] = field(default_factory=list)
    overloaded: bool = False
    runtime: type | None = None


@dataclass
class Members:
    consts: dict[str, str] = field(default_factory=dict)  # name -> "ann = value"
    props: dict[str, Prop] = field(default_factory=dict)
    methods: dict[str, Method] = field(default_factory=dict)
    order: list[str] = field(default_factory=list)

    def seen(self, name: str) -> bool:
        return name in self.consts or name in self.props or name in self.methods

    def add(self, name: str) -> None:
        if name not in self.order:
            self.order.append(name)


# --------------------------------------------------------------------------- helpers
def deco_names(fn: ast.FunctionDef) -> list[str]:
    out = []
    for d in fn.decorator_list:
        if isinstance(d, ast.Call):
            d = d.func
        out.append(ast.unparse(d))
    return out


def deprecated_decorator(fn: ast.FunctionDef) -> str | None:
    """Return the ``deprecated(...)`` decorator source of *fn*, if any."""
    for d in fn.decorator_list:
        target = d.func if isinstance(d, ast.Call) else d
        if ast.unparse(target).split(".")[-1] == "deprecated":
            src = ast.unparse(d)
            return "deprecated" + src[src.index("(") :] if "(" in src else "deprecated"
    return None


def load_runtime(path: str) -> type:
    mod, _, name = path.partition(":")
    return getattr(importlib.import_module(mod), name)


def split_union(ann: str) -> list[str]:
    """Split ``A | B[C | D]`` into ``["A", "B[C | D]"]`` (top-level ``|`` only)."""
    parts: list[str] = []
    depth = 0
    current: list[str] = []
    for ch in ann:
        if ch in "[(":
            depth += 1
        elif ch in "])":
            depth -= 1
        if ch == "|" and depth == 0:
            parts.append("".join(current).strip())
            current = []
        else:
            current.append(ch)
    parts.append("".join(current).strip())
    return [p for p in parts if p]


def merge_ann(a: str, b: str) -> str:
    """Union of two annotations (used when SPEC classes disagree about a property)."""
    if a == b:
        return a
    parts = split_union(a)
    for p in split_union(b):
        if p not in parts:
            parts.append(p)
    return " | ".join(parts)


def find_class(tree: ast.Module, cls_name: str, stub_file: str) -> ast.ClassDef:
    for node in tree.body:
        if isinstance(node, ast.ClassDef) and node.name == cls_name:
            return node
    sys.exit(
        f"error: class {cls_name!r} not found in stub {stub_file!r}; "
        "types-openpyxl may have been restructured -- update SPECS in tools/gen_proxies.py"
    )


# --------------------------------------------------------------------------- descriptors
# openpyxl's stubs annotate Serialisable attributes with the *descriptor* class
# (``pageMargins: Typed[PageMargins, Literal[True]]``).  On the raw class that
# works through ``__get__``/``__set__`` overloads; on a forwarding property it
# would make ``cs.pageMargins`` a descriptor object.  Map each descriptor to
# the value type it produces (``_N`` = allow_none).
_TYPED_LIKE = {
    "Typed",
    "Convertible",
    "Min",
    "Max",
    "MinMax",
    "Default",
    "TextPoint",
    "MatchPattern",
}
_VALUE_TYPED = {  # descriptor name -> fixed value type (``_N`` is the only parameter)
    "Integer": "int",
    "Float": "float",
    "Bool": "bool",
    "String": "str",
    "Text": "str",
    "ASCII": "bytes",
    "Tuple": "tuple[Any, ...]",
    "DateTime": "datetime",
}
_UNWRAP_FIRST = {"Descriptor", "Set", "Length", "Sequence"}  # X[T] -> T
_ALWAYS_ANY = {"Alias"}


def descriptor_aliases(tree: ast.Module) -> dict[str, str]:
    """Names a stub file imports from ``openpyxl.descriptors*`` -> descriptor name."""
    out: dict[str, str] = {}
    for node in tree.body:
        if isinstance(node, ast.ImportFrom) and (node.module or "").startswith(
            "openpyxl.descriptors"
        ):
            for alias in node.names:
                out[alias.asname or alias.name] = alias.name
    return out


def _allow_none(n: ast.expr | None) -> bool:
    # ``Literal[True]`` -> optional, ``Literal[False]`` -> required; anything
    # else (a TypeVar, missing) is treated as optional, the safe reading.
    if isinstance(n, ast.Subscript) and ast.unparse(n.value).split(".")[-1] == "Literal":
        return not (isinstance(n.slice, ast.Constant) and n.slice.value is False)
    return True


class _Expand(ast.NodeTransformer):
    def __init__(self, aliases: dict[str, str]) -> None:
        self.aliases = aliases

    def _descriptor(self, node: ast.expr) -> str | None:
        if isinstance(node, ast.Name):
            return self.aliases.get(node.id)
        return None

    def visit_Name(self, node: ast.Name) -> ast.expr:
        if self.aliases.get(node.id) in _ALWAYS_ANY:
            return ast.Name("Any", ast.Load())
        return node

    def visit_Subscript(self, node: ast.Subscript) -> ast.expr:
        self.generic_visit(node)
        kind = self._descriptor(node.value)
        if kind is None:
            return node
        args = list(node.slice.elts) if isinstance(node.slice, ast.Tuple) else [node.slice]
        result: ast.expr | None = None
        optional = False
        if kind in _TYPED_LIKE:
            result = args[0]
            optional = _allow_none(args[1] if len(args) > 1 else None)
        elif kind in _VALUE_TYPED:
            result = ast.parse(_VALUE_TYPED[kind], mode="eval").body
            optional = _allow_none(args[0] if args else None)
        elif kind == "NoneSet":
            result, optional = args[0], True
        elif kind in _UNWRAP_FIRST:
            result = args[0]
        if result is None:
            return node
        if optional:
            return ast.BinOp(result, ast.BitOr(), ast.Constant(None))
        return result


def expand_descriptors(ann: str, aliases: dict[str, str]) -> str:
    """Rewrite descriptor annotations in *ann* to the value types they produce."""
    if not aliases:
        return ann
    tree = ast.parse(ann, mode="eval")
    new = _Expand(aliases).visit(tree)
    return ast.unparse(ast.fix_missing_locations(new))


# --------------------------------------------------------------------------- collection
def collect(
    stubs: Path, specs: list[tuple[str, str, str]], exclude: set[str]
) -> tuple[Members, list[str]]:
    members = Members()
    imports: list[str] = []
    for stub_file, cls_name, runtime_path in specs:
        tree = ast.parse((stubs / stub_file).read_text())
        runtime = load_runtime(runtime_path)
        for node in tree.body:
            if isinstance(node, (ast.Import, ast.ImportFrom)):
                imports.append(ast.unparse(node))
        cls = find_class(tree, cls_name, stub_file)
        descriptors = descriptor_aliases(tree)
        for node in cls.body:
            if isinstance(node, ast.AnnAssign) and isinstance(node.target, ast.Name):
                name = node.target.id
                if name.startswith("_") or name in exclude or name in members.methods:
                    continue
                ann = expand_descriptors(ast.unparse(node.annotation), descriptors)
                if node.value is not None:
                    if not members.seen(name):
                        members.consts[name] = f"{ann} = {ast.unparse(node.value)}"
                        members.add(name)
                    continue
                if name in members.consts:
                    continue
                if ann.startswith("ClassVar["):
                    ann = ann[len("ClassVar[") : -1]
                prop = members.props.get(name)
                if prop is None:
                    members.props[name] = Prop(name, ann, ann)
                    members.add(name)
                else:
                    prop.getter = merge_ann(prop.getter, ann)
                    prop.setter = ann if prop.setter is None else merge_ann(prop.setter, ann)
            elif isinstance(node, ast.FunctionDef):
                name = node.name
                if name.startswith("_") or name in exclude:
                    continue
                decos = deco_names(node)
                if "classmethod" in decos or "staticmethod" in decos:
                    continue
                if "property" in decos:
                    if name in members.methods or name in members.consts:
                        continue
                    ret = ast.unparse(node.returns) if node.returns else "Any"
                    prop = members.props.get(name)
                    if prop is None:
                        members.props[name] = Prop(name, ret, deprecated=deprecated_decorator(node))
                        members.add(name)
                    else:
                        prop.getter = merge_ann(prop.getter, ret)
                        prop.deprecated = prop.deprecated or deprecated_decorator(node)
                elif any(d.endswith(".setter") for d in decos):
                    prop = members.props.get(name)
                    if prop is None:
                        continue
                    args = node.args.posonlyargs + node.args.args
                    if len(args) < 2:
                        warn(f"{cls_name}.{name}.setter has no value parameter in stub; skipped")
                        continue
                    arg = args[1]
                    ann = ast.unparse(arg.annotation) if arg.annotation else "Any"
                    prop.setter = ann if prop.setter is None else merge_ann(prop.setter, ann)
                else:
                    if name in members.props or name in members.consts:
                        continue
                    m = members.methods.get(name)
                    if m is None:
                        m = Method(name, runtime=runtime)
                        members.methods[name] = m
                        members.add(name)
                    elif m.runtime is not runtime:
                        continue  # already fully defined by an earlier class
                    if "overload" in decos:
                        m.overloaded = True
                    m.defs.append(node)
    return members, imports


def merge_imports(lines: list[str]) -> list[str]:
    """Merge ``from X import a`` / ``from X import a, b`` into one line per module."""
    plain: set[str] = set()
    by_module: dict[str, set[tuple[str, str | None]]] = {}
    for line in lines:
        node = ast.parse(line).body[0]
        if isinstance(node, ast.Import):
            plain.add(line)
            continue
        assert isinstance(node, ast.ImportFrom)
        module = "." * node.level + (node.module or "")
        names = by_module.setdefault(module, set())
        drop = RUNTIME_IMPORTS.get(module, set())
        for alias in node.names:
            if alias.name in drop and alias.asname is None:
                continue
            names.add((alias.name, alias.asname))
    out = sorted(plain)
    for module in sorted(by_module):
        names = by_module[module]
        if not names:
            continue
        rendered = ", ".join(
            n if a is None else f"{n} as {a}"
            for n, a in sorted(names, key=lambda na: (na[0], na[1] or ""))
        )
        out.append(f"from {module} import {rendered}")
    return out


# --------------------------------------------------------------------------- rendering
def runtime_signature(m: Method) -> inspect.Signature | None:
    fn = inspect.getattr_static(m.runtime, m.name, None)
    if fn is None:
        return None
    try:
        return inspect.signature(fn)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None


def runtime_default(m: Method, pname: str) -> str:
    sig = runtime_signature(m)
    p = sig.parameters.get(pname) if sig is not None else None
    if p is None or p.default is inspect.Parameter.empty:
        warn(
            f"{m.runtime.__name__ if m.runtime else '?'}.{m.name}({pname}=...): "
            "stub default is '...' but the runtime default could not be resolved; using None"
        )
        return "None"
    return repr(p.default)


def render_signature(fn: ast.FunctionDef, m: Method) -> tuple[str, str]:
    """Return (parameter list source, forwarding call arguments source)."""
    a = fn.args
    params: list[str] = ["self"]
    call: list[str] = []

    def ann(arg: ast.arg) -> str:
        return f": {ast.unparse(arg.annotation)}" if arg.annotation else ""

    def default(arg_default: ast.expr | None, pname: str) -> str:
        if arg_default is None:
            return ""
        if isinstance(arg_default, ast.Constant) and arg_default.value is Ellipsis:
            return f" = {runtime_default(m, pname)}"
        return f" = {ast.unparse(arg_default)}"

    positional = a.posonlyargs + a.args
    n_defaults = len(a.defaults)
    defaults = [None] * (len(positional) - n_defaults) + list(a.defaults)
    for i, arg in enumerate(positional):
        if arg.arg == "self":
            continue
        params.append(f"{arg.arg}{ann(arg)}{default(defaults[i], arg.arg)}")
        if arg in a.posonlyargs:
            call.append(arg.arg)
        else:
            call.append(f"{arg.arg}={arg.arg}")
    posonly = [p for p in a.posonlyargs if p.arg != "self"]
    if posonly:
        params.insert(len(posonly) + 1, "/")
    if a.vararg:
        params.append(f"*{a.vararg.arg}{ann(a.vararg)}")
        call.append(f"*{a.vararg.arg}")
    elif a.kwonlyargs:
        params.append("*")
    for arg, d in zip(a.kwonlyargs, a.kw_defaults, strict=True):
        params.append(f"{arg.arg}{ann(arg)}{default(d, arg.arg)}")
        call.append(f"{arg.arg}={arg.arg}")
    if a.kwarg:
        params.append(f"**{a.kwarg.arg}{ann(a.kwarg)}")
        call.append(f"**{a.kwarg.arg}")
    return ", ".join(params), ", ".join(call)


def render_property(p: Prop) -> list[str]:
    lines = ["    @property"]
    if p.deprecated:
        lines.append(f"    @{p.deprecated}")
    lines.append(f"    def {p.name}(self) -> {p.getter}:")
    if p.getter in NEVER_TYPES:
        lines += [
            f"        self._raw.{p.name}",
            f'        raise AssertionError("openpyxl declares {p.name} as {p.getter}")',
        ]
    else:
        lines.append(f"        return self._raw.{p.name}")
    lines.append("")
    if p.setter is not None:
        lines += [
            f"    @{p.name}.setter",
            f"    def {p.name}(self, value: {p.setter}) -> None:",
            f"        self._raw.{p.name} = value",
            "",
        ]
    return lines


def render_method(m: Method, overlapping: bool) -> list[str]:
    lines: list[str] = []
    name = m.name
    if m.overloaded:
        for i, fn in enumerate(m.defs):
            params, _ = render_signature(fn, m)
            ret = ast.unparse(fn.returns) if fn.returns else "Any"
            dep = deprecated_decorator(fn)
            lines.append("    @overload")
            if dep:
                lines.append(f"    @{dep}")
            tail = ""
            if overlapping and i == 0:
                tail = "  # pyright: ignore[reportOverlappingOverload]"
            lines += [f"    def {name}({params}) -> {ret}: ...{tail}", ""]
        lines += [
            f"    def {name}(self, *args: Any, **kwargs: Any) -> Any:",
            f"        return self._raw.{name}(*args, **kwargs)",
            "",
        ]
        return lines
    fn = m.defs[0]
    params, call = render_signature(fn, m)
    ret = ast.unparse(fn.returns) if fn.returns else "Any"
    dep = deprecated_decorator(fn)
    if dep:
        lines.append(f"    @{dep}")
    lines.append(f"    def {name}({params}) -> {ret}:")
    if ret in NEVER_TYPES:
        lines += [
            f"        self._raw.{name}({call})",
            f'        raise AssertionError("openpyxl declares {name} as {ret}")',
        ]
    else:
        lines.append(f"        return self._raw.{name}({call})")
    lines.append("")
    return lines


def render_class(
    cls_name: str,
    members: Members,
    extra: dict[str, tuple[str, str | None]],
    raw_type: str,
    imported: set[str],
) -> str:
    lines = [
        f"class {cls_name}(GuardedProxy):",
        f'    """Typed synchronous forwarders generated from types-openpyxl ({raw_type})."""',
        "",
        "    __slots__ = ()",
        "",
    ]
    for name in members.order:
        if name in members.consts:
            lines.append(f"    {name}: {members.consts[name]}")
    lines.append("")
    overlapping = OVERLAPPING_OVERLOADS.get(cls_name, set())
    for name in members.order:
        if name in members.props:
            p = members.props[name]
            if name in imported:  # the property would shadow the type name in annotations
                p = Prop(name, "Any", "Any" if p.setter is not None else None, p.deprecated)
            lines += render_property(p)
        elif name in members.methods:
            lines += render_method(members.methods[name], name in overlapping)
    for name, (ann, _imp) in extra.items():
        lines += render_property(Prop(name, ann, ann))
    return "\n".join(lines)


# --------------------------------------------------------------------------- self-checks
def check_signature_drift(cls_name: str, members: Members) -> list[str]:
    """Compare stub parameter names/defaults of plain forwarders with the runtime."""
    problems: list[str] = []
    for m in members.methods.values():
        if m.overloaded or not m.defs:
            continue
        sig = runtime_signature(m)
        where = f"{cls_name}.{m.name} ({m.runtime.__module__}.{m.runtime.__name__})"  # type: ignore[union-attr]
        if sig is None:
            warn(f"{where}: runtime signature unavailable; drift check skipped")
            continue
        fn = m.defs[0]
        a = fn.args
        stub: list[tuple[str, str | None]] = []
        positional = a.posonlyargs + a.args
        defaults = [None] * (len(positional) - len(a.defaults)) + list(a.defaults)
        for arg, d in zip(positional, defaults, strict=True):
            if arg.arg == "self":
                continue
            stub.append((arg.arg, None if d is None else ast.unparse(d)))
        if a.vararg:
            stub.append((a.vararg.arg, None))
        for arg, d in zip(a.kwonlyargs, a.kw_defaults, strict=True):
            stub.append((arg.arg, None if d is None else ast.unparse(d)))
        if a.kwarg:
            stub.append((a.kwarg.arg, None))

        runtime: list[tuple[str, str | None]] = []
        for p in sig.parameters.values():
            if p.name == "self":
                continue
            has_default = p.default is not inspect.Parameter.empty
            runtime.append((p.name, repr(p.default) if has_default else None))

        if [n for n, _ in stub] != [n for n, _ in runtime]:
            problems.append(
                f"{where}: parameter names differ\n"
                f"    stub:    {[n for n, _ in stub]}\n"
                f"    runtime: {[n for n, _ in runtime]}"
            )
            continue
        for (name, sd), (_, rd) in zip(stub, runtime, strict=True):
            if sd == "...":
                continue  # "use the runtime default" -- resolved by runtime_default()
            if (sd is None) != (rd is None) or (sd is not None and sd != rd):
                problems.append(f"{where}: default of {name!r} differs: stub={sd} runtime={rd}")
    return problems


def runtime_instances() -> dict[str, list[Any]]:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    cs = wb.create_chartsheet()
    wo = openpyxl.Workbook(write_only=True).create_sheet()
    buf = io.BytesIO()
    openpyxl.Workbook().save(buf)
    buf.seek(0)
    ro = openpyxl.load_workbook(buf, read_only=True).active
    return {
        "WorkbookProxy": [wb],
        "WorksheetProxy": [ws, ro, wo],
        "ChartsheetProxy": [cs],
    }


def check_missing_members(cls_name: str, emitted: set[str], instances: list[Any]) -> None:
    """Warn about public runtime attributes the generated class does not define."""
    known = emitted | EXCLUDE[cls_name]
    for inst in instances:
        missing = []
        for name in dir(inst):
            if name.startswith("_") or name in known:
                continue
            static = inspect.getattr_static(type(inst), name, None)
            if isinstance(static, (classmethod, staticmethod)):
                continue
            missing.append(name)
        if missing:
            warn(
                f"{cls_name}: {type(inst).__name__} has public members unknown to the stubs, "
                f"EXCLUDE and EXTRA_MEMBERS: {sorted(missing)}"
            )


# --------------------------------------------------------------------------- module shims
SHIM_HEADER = (
    "# Generated by tools/gen_proxies.py -- do not edit.\n"
    "# Re-exports the types-openpyxl stub (typeshed, Apache-2.0) for the openpyxl module\n"
    "# aliased at this path; see _compat.py.\n"
)


def generate_module_shims(stubs: Path) -> tuple[int, int, int]:
    """Write ``.pyi`` shims for every public openpyxl module that has a stub.

    Returns (written or kept, removed, skipped-because-of-own-module).
    """
    import openpyxl

    wanted: dict[Path, str] = {}
    skipped_own = 0
    for info in pkgutil.walk_packages(openpyxl.__path__, "openpyxl."):
        parts = info.name.split(".")[1:]
        if any(p.startswith("_") for p in parts):
            continue
        rel = Path(*parts)
        if info.ispkg:
            stub = stubs / rel / "__init__.pyi"
            target = PKG_DIR / rel / "__init__.pyi"
        else:
            stub = (stubs / rel).with_suffix(".pyi")
            target = (PKG_DIR / rel).with_suffix(".pyi")
        if not stub.is_file():
            continue
        own_module = (PKG_DIR / rel).with_suffix(".py")
        own_package = PKG_DIR / rel / "__init__.py"
        if own_module.exists() or own_package.exists():
            warn(f"shim for {info.name} skipped: aioopenpyxl defines that module itself")
            skipped_own += 1
            continue
        wanted[target] = f"{SHIM_HEADER}from {info.name} import *\n"

    existing = {p for p in PKG_DIR.rglob("*.pyi") if "__pycache__" not in p.parts}
    removed = 0
    for stale in sorted(existing - set(wanted)):
        stale.unlink()
        removed += 1
    for path, text in wanted.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        if not path.exists() or path.read_text() != text:
            path.write_text(text)
    # Drop directories that only held shims which no longer exist.
    for d in sorted((p for p in PKG_DIR.rglob("*") if p.is_dir()), reverse=True):
        if d.name != "__pycache__" and not any(d.iterdir()):
            d.rmdir()
    return len(wanted), removed, skipped_own


# --------------------------------------------------------------------------- main
def main() -> None:
    stubs = stub_dir()
    import openpyxl

    types_version = importlib.metadata.version("types-openpyxl")

    all_imports: list[str] = []
    collected: list[tuple[str, Members]] = []
    for cls_name, specs in SPECS.items():
        members, imports = collect(stubs, specs, EXCLUDE[cls_name])
        all_imports += imports
        for _name, (_ann, imp) in EXTRA_MEMBERS[cls_name].items():
            if imp:
                all_imports.append(imp)
        collected.append((cls_name, members))

    drift: list[str] = []
    instances = runtime_instances()
    for cls_name, members in collected:
        drift += check_signature_drift(cls_name, members)
        emitted = set(members.order) | set(EXTRA_MEMBERS[cls_name])
        check_missing_members(cls_name, emitted, instances[cls_name])
    if drift:
        print("error: stub/runtime signature drift detected:", file=sys.stderr)
        for p in drift:
            print(f"  {p}", file=sys.stderr)
        sys.exit(1)

    merged_imports = merge_imports(all_imports)
    imported_names: set[str] = set()
    for line in merged_imports:
        for alias in ast.parse(line).body[0].names:  # type: ignore[attr-defined]
            imported_names.add(alias.asname or alias.name)
    classes = [
        render_class(c, m, EXTRA_MEMBERS[c], RAW_TYPES[c], imported_names) for c, m in collected
    ]

    header = f'''"""Typed synchronous forwarders for the complete openpyxl surface.

GENERATED FILE -- do not edit.  Regenerate with ``uv run python tools/gen_proxies.py``.
Source: types-openpyxl {types_version} stubs for openpyxl {openpyxl.__version__}.
The signatures below are derived from the types-openpyxl stubs maintained in typeshed
(https://github.com/python/typeshed, Apache-2.0); openpyxl itself is MIT-licensed and is
used as a dependency, not vendored.  aioopenpyxl is not affiliated with either project.
"""
# ruff: noqa
# fmt: off

from __future__ import annotations

from typing import TYPE_CHECKING, Any, overload

from typing_extensions import deprecated

from ._base import GuardedProxy

if TYPE_CHECKING:
'''
    body = (
        header
        + "\n".join(f"    {line}" for line in merged_imports)
        + "\n\n\n"
        + "\n\n".join(classes)
    )
    body = body.rstrip("\n") + "\n"  # what ``ruff format`` would produce
    OUT.write_text(body)
    print(f"wrote {OUT.relative_to(ROOT)} ({len(body.splitlines())} lines)")

    kept, removed, skipped = generate_module_shims(stubs)
    print(
        f"module shims: {kept} .pyi files under {PKG_DIR.relative_to(ROOT)} "
        f"({removed} stale removed, {skipped} skipped)"
    )


if __name__ == "__main__":
    main()
