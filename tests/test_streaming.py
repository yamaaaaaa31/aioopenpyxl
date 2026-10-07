"""``iter_rows`` / ``iter_cols`` / ``read_rows``, chunks, ``prefetch`` and clean-up."""

from __future__ import annotations

import contextlib
import gc
import subprocess
import sys
import threading
from pathlib import Path
from typing import Any

import anyio
import anyio.lowlevel
import anyio.to_thread
import openpyxl
import pytest
from openpyxl.cell.read_only import ReadOnlyCell

import aioopenpyxl
from aioopenpyxl import WorkbookBusyError
from tests.conftest import ROWS, abandons_generator, no_errors_logged, prefetch_threads

pytestmark = pytest.mark.anyio


# -------------------------------------------------------------- basic chunks
@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize("chunk_size", [1, 7, 1024])
async def test_iter_rows_values_only_in_chunks(
    sample_xlsx: Path, read_only: bool, chunk_size: int
) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=read_only) as wb:
        ws = wb["Data"]
        assert ws.is_read_only is read_only
        rows = [row async for row in ws.iter_rows(values_only=True, chunk_size=chunk_size)]
    assert rows == ROWS


async def test_iter_rows_with_bounds_yields_cells(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        ws = wb["Data"]
        rows = [row async for row in ws.iter_rows(min_row=2, max_row=3, max_col=2)]
    assert [[c.value for c in row] for row in rows] == [[1, "user1"], [2, "user2"]]
    first = rows[0][0]
    assert isinstance(first, ReadOnlyCell) and first.coordinate == "A2"


async def test_rows_values_and_columns_properties(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx) as wb:
        ws = wb["Data"]
        n_rows = sum([1 async for _ in ws.rows])
        values = [row async for row in ws.values]
        cols = [tuple(c.value for c in col) async for col in ws.columns]
    assert n_rows == len(ROWS)
    assert values == ROWS
    assert cols[0] == tuple(r[0] for r in ROWS)


async def test_iter_rows_rejects_bad_chunk_size(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx) as wb:
        with pytest.raises(ValueError, match="chunk_size"):
            async for _ in wb["Data"].iter_rows(chunk_size=0):
                pass


async def test_breaking_out_of_iteration_closes_generator(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        ws = wb["Data"]
        agen = ws.iter_rows(values_only=True, chunk_size=3)
        first = await anext(agen)
        await agen.aclose()
    assert first == ROWS[0]


async def test_read_rows_single_hop(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        ws = wb["Data"]
        assert await ws.read_rows(min_row=1, max_row=2) == ROWS[:2]
        cells = await ws.read_rows(max_row=1, values_only=False)
        assert [c.value for c in cells[0]] == list(ROWS[0])


async def test_sync_access_between_iteration_chunks_is_allowed() -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    await ws.append_rows([[i] for i in range(10)])
    seen = 0
    async for _row in ws.iter_rows(chunk_size=3):
        assert ws.max_row == 10  # lock is released between chunk fetches
        seen += 1
    assert seen == 10
    assert [v async for (v,) in ws.values] == list(range(10))
    assert sum([1 async for _ in ws]) == 10


@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize("prefetch", [0, 1, 2, 3])
@pytest.mark.parametrize("chunk_size", [3, 5, 17, 51, 100])  # ROWS has 51 rows
async def test_every_chunk_size_yields_all_rows(
    sample_xlsx: Path, read_only: bool, prefetch: int, chunk_size: int
) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=read_only) as wb:
        ws = wb["Data"]
        rows = [
            r
            async for r in ws.iter_rows(values_only=True, chunk_size=chunk_size, prefetch=prefetch)
        ]
        assert rows == ROWS
        # A short final chunk ends the stream without fetching an empty tail,
        # and in every case the runner is released once the loop is over.
        assert wb._runner.busy is False
        assert prefetch_threads() == []


@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize("prefetch", [0, 2])
async def test_empty_range_terminates(sample_xlsx: Path, read_only: bool, prefetch: int) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=read_only) as wb:
        ws = wb["Data"]
        rows = [r async for r in ws.iter_rows(min_row=5, max_row=4, prefetch=prefetch)]
        assert rows == []
        assert wb._runner.busy is False
        if read_only:
            # An empty read-only sheet has no rows at all.
            assert [r async for r in wb["Empty"].iter_rows(prefetch=prefetch)] == []
            assert wb._runner.busy is False


# ------------------------------------------------------------------ prefetch
@pytest.mark.parametrize("read_only", [False, True])
async def test_prefetch_streams_all_rows(sample_xlsx: Path, read_only: bool) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=read_only) as wb:
        ws = wb["Data"]
        rows = [row async for row in ws.iter_rows(values_only=True, chunk_size=7, prefetch=2)]
        assert rows == ROWS
        # The producer thread has stopped: the sheet is usable again.
        assert wb._runner.busy is False
        assert await ws.read_rows(max_row=1) == ROWS[:1]


async def test_prefetch_marks_the_sheet_busy_during_iteration(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx) as wb:
        ws = wb["Data"]
        n = 0
        async for _ in ws.iter_rows(chunk_size=5, prefetch=1):
            n += 1
            if n == 3:
                with pytest.raises(WorkbookBusyError):
                    _ = ws.max_row
        assert n == len(ROWS)
        assert ws.max_row == len(ROWS)


async def test_prefetch_early_exit_stops_the_producer(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        ws = wb["Data"]
        agen = ws.iter_rows(values_only=True, chunk_size=2, prefetch=3)
        assert await anext(agen) == ROWS[0]
        await agen.aclose()
        assert wb._runner.busy is False
        # the raw sheet is intact
        assert await ws.read_rows(max_row=2) == ROWS[:2]


async def test_prefetch_propagates_producer_errors(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx) as wb:
        ws = wb["Data"]

        def boom():
            raise RuntimeError("parse failed")
            yield  # pragma: no cover - makes this a generator

        with pytest.raises(RuntimeError, match="parse failed"):
            async for _ in ws._stream(boom, 4, 2):
                pass
        assert wb._runner.busy is False

        with pytest.raises(RuntimeError, match="parse failed"):
            async for _ in ws._stream(boom, 4, 0):
                pass
        assert wb._runner.busy is False


async def test_invalid_prefetch_rejected(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx) as wb:
        with pytest.raises(ValueError, match="prefetch"):
            async for _ in wb["Data"].iter_rows(prefetch=-1):
                pass


async def test_iter_cols_with_prefetch(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx) as wb:
        ws = wb["Data"]
        cols = [c async for c in ws.iter_cols(values_only=True, chunk_size=2, prefetch=1)]
        assert cols == [tuple(r[i] for r in ROWS) for i in range(3)]
        assert wb._runner.busy is False


# ------------------------------------------------------------ concurrent reads
@pytest.mark.parametrize("read_only", [False, True])
async def test_concurrent_read_rows_serialise_instead_of_failing(
    sample_xlsx: Path, read_only: bool
) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=read_only) as wb:
        ws = wb["Data"]
        results: list[Any] = []

        async def read() -> None:
            results.append(await ws.read_rows())

        async with anyio.create_task_group() as tg:
            for _ in range(3):
                tg.start_soon(read)
        assert results == [ROWS, ROWS, ROWS]
        assert wb._runner.busy is False


@pytest.mark.parametrize("read_only", [False, True])
async def test_concurrent_iter_rows_serialise_instead_of_failing(
    sample_xlsx: Path, read_only: bool
) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=read_only) as wb:
        ws = wb["Data"]
        results: dict[int, list[Any]] = {}

        async def consume(prefetch: int) -> None:
            results[prefetch] = [
                r async for r in ws.iter_rows(values_only=True, chunk_size=7, prefetch=prefetch)
            ]

        async def read() -> None:
            assert await ws.read_rows(max_row=1) == ROWS[:1]

        async with anyio.create_task_group() as tg:
            tg.start_soon(consume, 0)
            tg.start_soon(consume, 2)
            tg.start_soon(read)
            tg.start_soon(consume, 1)
        assert results == {0: ROWS, 1: ROWS, 2: ROWS}
        assert wb._runner.busy is False


async def test_iter_rows_while_a_thread_owns_the_workbook_waits() -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    await ws.append_rows([[i] for i in range(10)])
    started = threading.Event()
    release = threading.Event()

    def slow(raw: openpyxl.Workbook) -> None:
        started.set()
        release.wait(5)

    seen: list[Any] = []

    async def consume() -> None:
        seen.extend([r async for r in ws.iter_rows(values_only=True, chunk_size=4)])

    async with anyio.create_task_group() as tg:
        tg.start_soon(wb.run, slow)
        await anyio.to_thread.run_sync(started.wait, 5)
        tg.start_soon(consume)
        await anyio.sleep(0.05)
        assert seen == []  # queued behind the running call, no WorkbookBusyError
        release.set()
    assert seen == [(i,) for i in range(10)]


async def test_write_only_sheet_iteration_fails_inside_the_thread() -> None:
    wb = aioopenpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("W")
    with pytest.raises(AttributeError):
        async for _ in ws.iter_rows():
            pass
    assert wb._runner.busy is False
    with pytest.raises(AttributeError):
        await ws.read_rows()


# ------------------------------------------------- abandoned loops and clean-up
@abandons_generator
@pytest.mark.parametrize("read_only", [False, True])
async def test_abandoned_prefetch_loop_releases_the_runner(
    sample_xlsx: Path, read_only: bool, caplog: pytest.LogCaptureFixture
) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=read_only) as wb:
        ws = wb["Data"]
        async for _ in ws.iter_rows(prefetch=2, chunk_size=4):
            break
        # The generator is now unreferenced; its aclose() runs from the event
        # loop's finaliser hook in *another* task.  Until then the workbook may
        # still read as busy, but awaiting the runner must not deadlock.
        gc.collect()
        await anyio.lowlevel.checkpoint()
        with anyio.fail_after(5):
            assert await wb.run(lambda raw: raw.sheetnames) == ["Data", "Empty"]
        assert wb._runner.busy is False
        assert prefetch_threads() == []
        assert await ws.read_rows(max_row=1) == ROWS[:1]
    # ``async with`` closed the workbook; closing again is harmless.
    with anyio.fail_after(5):
        await wb.close()
    assert no_errors_logged(caplog)


async def test_breaking_out_with_aclosing_prefetch_zero(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        ws = wb["Data"]
        seen = []
        async with contextlib.aclosing(ws.iter_rows(values_only=True, chunk_size=3)) as rows:
            async for row in rows:
                seen.append(row)
                if len(seen) == 2:
                    break
        assert seen == ROWS[:2]
        assert wb._runner.busy is False
        assert await ws.read_rows(max_row=1) == ROWS[:1]


@pytest.mark.parametrize("prefetch", [0, 2])
async def test_iterator_is_closed_in_a_worker_thread_under_the_runner(
    sample_xlsx: Path, prefetch: int
) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        ws = wb["Data"]
        raw: Any = ws.wrapped
        loop_thread = threading.get_ident()
        closed_on: dict[str, Any] = {}
        real = raw.iter_rows

        class Spy:
            def __init__(self, gen: Any) -> None:
                self.gen = gen

            def __iter__(self) -> Spy:
                return self

            def __next__(self) -> Any:
                return next(self.gen)

            def close(self) -> None:
                closed_on["thread"] = threading.get_ident()
                closed_on["busy"] = wb._runner.busy
                self.gen.close()

        raw.iter_rows = lambda **kw: Spy(real(**kw))

        rows = ws.iter_rows(chunk_size=2, prefetch=prefetch)
        await anext(rows)
        await rows.aclose()
        assert closed_on["thread"] != loop_thread
        assert closed_on["busy"] is True
        assert wb._runner.busy is False

        # A complete iteration leaves an exhausted generator: no close hop needed.
        closed_on.clear()
        assert len([r async for r in ws.iter_rows(chunk_size=20, prefetch=prefetch)]) == len(ROWS)
        assert closed_on == {}


async def test_stream_aclose_is_idempotent_and_waits_for_the_first_caller(
    sample_xlsx: Path,
) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        ws = wb["Data"]
        raw: Any = ws.wrapped
        for prefetch in (0, 2):
            stream = ws._stream(lambda: raw.iter_rows(values_only=True), 2, prefetch)
            assert await anext(stream) == ROWS[0]
            # Concurrent closers: the second waits for the first, nothing runs twice.
            async with anyio.create_task_group() as tg:
                tg.start_soon(stream.aclose)
                tg.start_soon(stream.aclose)
                tg.start_soon(stream.aclose)
            assert wb._runner.busy is False
            await stream.aclose()  # and again, sequentially
            assert [r async for r in stream] == []
        # The public generator is closed once by ``aclose()``; closing again is a no-op.
        rows = ws.iter_rows(chunk_size=2, prefetch=2)
        await anext(rows)
        await rows.aclose()
        await rows.aclose()
        assert wb._runner.busy is False


# ------------------------------------------------------- loop shutdown (asyncio)
_SHUTDOWN_SCRIPT = """
import asyncio, threading, warnings, aioopenpyxl
warnings.simplefilter("error")
keep = {}
async def main():
    wb = aioopenpyxl.Workbook(); ws = wb.active
    await ws.append_rows([[i] for i in range(50)])
    for prefetch in (0, 2):
        agen = ws.iter_rows(prefetch=prefetch, chunk_size=3)
        assert (await anext(agen))[0].value == 0
        keep[prefetch] = agen  # left open on purpose: shutdown_asyncgens() closes it
    keep["wb"] = wb
asyncio.run(main())
assert keep["wb"]._runner.busy is False
assert not [t for t in threading.enumerate() if t.name == "aioopenpyxl-prefetch"]
print("ok")
"""


def test_open_generators_are_closed_cleanly_at_asyncio_shutdown() -> None:
    proc = subprocess.run(
        [sys.executable, "-c", _SHUTDOWN_SCRIPT],
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert proc.returncode == 0, proc.stderr
    assert "already running" not in proc.stderr
    assert "error occurred during closing" not in proc.stderr
    assert proc.stdout.strip() == "ok"


async def test_prefetch_producer_sees_the_callers_context(sample_xlsx: Path) -> None:
    """The producer thread runs in a copy of the caller's context (like anyio workers)."""
    import contextvars

    var: contextvars.ContextVar[str] = contextvars.ContextVar("aioopenpyxl_test_var")
    var.set("request-42")
    seen: list[str] = []

    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        ws = wb["Data"]
        raw: Any = ws.wrapped
        real = raw.iter_rows

        def spying_iter_rows(**kwargs: Any) -> Any:
            seen.append(var.get("<missing>"))
            return real(**kwargs)

        raw.iter_rows = spying_iter_rows
        rows = [r async for r in ws.iter_rows(values_only=True, prefetch=2, chunk_size=7)]
    assert rows == ROWS
    assert seen == ["request-42"]
