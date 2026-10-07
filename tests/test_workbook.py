"""``Workbook`` API: loading, saving, sheet management, the class itself, chartsheets."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any, cast

import anyio
import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.chartsheet import Chartsheet
from openpyxl.worksheet._write_only import WriteOnlyWorksheet
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.xml.functions import Element, fromstring

import aioopenpyxl
from aioopenpyxl import AsyncChartsheet, AsyncWorkbook, AsyncWorksheet, Workbook
from tests.conftest import ROWS

pytestmark = pytest.mark.anyio


async def test_load_workbook_is_awaitable(sample_xlsx: Path) -> None:
    wb = await aioopenpyxl.load_workbook(sample_xlsx)
    try:
        assert isinstance(wb, AsyncWorkbook)
        assert wb.sheetnames == ["Data", "Empty"]
        assert isinstance(wb.wrapped, openpyxl.Workbook)
    finally:
        await wb.close()


async def test_load_workbook_as_context_manager_closes_archive(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        assert wb.read_only is True  # proxied attribute
        archive = cast(Any, wb.wrapped)._archive  # private openpyxl state
        assert archive.fp is not None
    assert archive.fp is None  # zipfile closed by wb.close()


async def test_aexit_closes_the_vba_archive_copy(sample_xlsx: Path) -> None:
    # ``keep_vba=True`` makes openpyxl keep an append-mode ZipFile over a
    # BytesIO that it never closes; its finaliser can then fail noisily at
    # interpreter shutdown.  Leaving the block closes it.
    async with aioopenpyxl.load_workbook(sample_xlsx, keep_vba=True) as wb:
        vba_archive = wb.wrapped.vba_archive
        assert vba_archive is not None and vba_archive.fp is not None
    assert vba_archive.fp is None


async def test_close_keeps_the_vba_archive_like_openpyxl(sample_xlsx: Path) -> None:
    # ``close()`` alone is exactly ``openpyxl.Workbook.close``: the copy stays
    # readable, so a save after ``close()`` still works as it does in openpyxl.
    wb = await aioopenpyxl.load_workbook(sample_xlsx, keep_vba=True)
    await wb.close()
    vba_archive = wb.wrapped.vba_archive
    assert vba_archive is not None and vba_archive.fp is not None
    assert len(await wb.to_bytes()) > 0
    async with wb:
        pass
    assert vba_archive.fp is None


async def test_aexit_without_vba_archive_is_unchanged(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx) as wb:
        assert wb.wrapped.vba_archive is None
    async with aioopenpyxl.Workbook() as created:
        assert created.wrapped.vba_archive is None


def test_subscripted_wrappers_construct_at_runtime() -> None:
    # ``Workbook[Worksheet]()`` goes through ``typing``'s generic alias, which
    # tries to set ``__orig_class__`` on the slotted, ``__setattr__``-guarded
    # instance; that must stay harmless.
    wb = Workbook[Worksheet]()
    assert isinstance(wb, Workbook)
    assert type(wb) is Workbook
    ws = wb.active
    assert isinstance(ws, AsyncWorksheet)
    write_only = Workbook[WriteOnlyWorksheet](write_only=True)
    assert write_only.write_only is True
    wrapped = Workbook[Worksheet].wrap(openpyxl.Workbook())
    assert isinstance(wrapped, Workbook)
    raw_ws = openpyxl.Workbook().active
    assert raw_ws is not None
    assert AsyncWorksheet[Worksheet](raw_ws).wrapped is raw_ws
    assert aioopenpyxl.AsyncWorkbook[Worksheet] is not None


async def test_workbook_sheet_management() -> None:
    wb = aioopenpyxl.Workbook()
    assert len(wb) == 1
    first = wb.active
    assert isinstance(first, AsyncWorksheet)

    second = wb.create_sheet("Second")
    assert wb.sheetnames == ["Sheet", "Second"]
    assert "Second" in wb
    assert wb["Second"] == second
    assert wb.index(second) == 1
    assert [ws.title for ws in wb] == ["Sheet", "Second"]

    wb.active = second
    assert wb.active == second
    wb.active = 0
    assert wb.active == first

    copy = wb.copy_worksheet(second)
    assert copy.title == "Second Copy"
    wb.move_sheet(copy, offset=-2)
    assert wb.sheetnames[0] == "Second Copy"

    wb.remove(copy)
    del wb["Second"]
    assert wb.sheetnames == ["Sheet"]


async def test_wrappers_are_cached_per_raw_sheet() -> None:
    wb = aioopenpyxl.Workbook()
    assert wb.active is wb.active
    ws = wb.create_sheet("S")
    assert wb["S"] is ws
    assert wb.worksheets[1] is ws
    assert ws.parent is wb
    cs = wb.create_chartsheet("C")
    assert wb.chartsheets[0] is cs


async def test_workbook_setattr_proxies_to_wrapped() -> None:
    wb = aioopenpyxl.Workbook()
    wb.iso_dates = True
    assert wb.wrapped.iso_dates is True


async def test_save_and_reload_round_trip(tmp_path: Path) -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    appended = await ws.append_rows(ROWS)
    assert appended == len(ROWS)
    out = tmp_path / "out.xlsx"
    await wb.save(out)

    async with aioopenpyxl.load_workbook(out, read_only=True) as loaded:
        rows = [row async for row in loaded["Data"].values]
    assert rows == ROWS


async def test_save_to_binary_stream() -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    await ws.append(["x"])
    buf = io.BytesIO()
    await wb.save(buf)
    buf.seek(0)
    raw = openpyxl.load_workbook(buf).active
    assert raw is not None
    assert raw["A1"].value == "x"


async def test_write_only_workbook(tmp_path: Path) -> None:
    wb = aioopenpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("W")
    assert ws.is_write_only
    await ws.append(["a", "b"])
    await ws.append_rows([[1, 2], [3, 4]])
    out = tmp_path / "wo.xlsx"
    await wb.save(out)

    async with aioopenpyxl.load_workbook(out) as loaded:
        assert await loaded["W"].read_rows() == [("a", "b"), (1, 2), (3, 4)]


async def test_run_executes_against_raw_workbook(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx) as wb:
        titles = await wb.run(lambda raw: [ws.title for ws in raw.worksheets])
    assert titles == ["Data", "Empty"]


async def test_custom_limiter_is_kept(sample_xlsx: Path) -> None:
    limiter = anyio.CapacityLimiter(1)
    async with aioopenpyxl.load_workbook(sample_xlsx, limiter=limiter) as wb:
        assert wb.limiter is limiter
        assert wb.active is not None and wb.active.parent is wb
    wb2 = aioopenpyxl.Workbook(limiter=limiter)
    assert wb2.limiter is limiter
    assert aioopenpyxl.Workbook().limiter is None


async def test_saving_read_only_workbook_raises(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        with pytest.raises(TypeError):
            await wb.save(io.BytesIO())


async def test_to_bytes_and_load_workbook_bytes_round_trip() -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    await ws.append_rows(ROWS)
    data = await wb.to_bytes()
    assert data[:2] == b"PK"

    async with aioopenpyxl.load_workbook_bytes(data, read_only=True) as loaded:
        active = loaded.active
        assert active is not None
        assert [r async for r in active.values] == ROWS
    loaded2 = await aioopenpyxl.load_workbook_bytes(bytearray(data))
    assert loaded2.sheetnames == ["Sheet"]
    await loaded2.close()


# ------------------------------------------------------------- Workbook class
def test_workbook_is_a_class_with_openpyxl_constructor_signature() -> None:
    wb = Workbook()
    assert isinstance(wb, Workbook)
    assert isinstance(wb, AsyncWorkbook)
    assert AsyncWorkbook is Workbook
    assert Workbook(write_only=True).write_only is True
    assert Workbook(iso_dates=True).iso_dates is True
    assert Workbook(True, True).write_only is True  # positional, like openpyxl
    # Outside any event loop: the lazily created semaphore makes this safe.
    assert wb.sheetnames == ["Sheet"]
    assert repr(wb) == "<Workbook sheets=['Sheet']>"


def test_workbook_wrap_keeps_the_raw_workbook() -> None:
    raw = openpyxl.Workbook()
    raw.create_sheet("Extra")
    limiter = anyio.CapacityLimiter(2)
    wb = Workbook.wrap(raw, limiter=limiter)
    assert isinstance(wb, Workbook)
    assert wb.wrapped is raw
    assert wb.limiter is limiter
    assert wb.sheetnames == ["Sheet", "Extra"]
    assert repr(wb) == "<Workbook sheets=['Sheet', 'Extra']>"
    assert Workbook.wrap(raw).wrapped is raw
    # A raw workbook has exactly one Runner: a second wrapper shares it (and
    # its limiter); asking for a different limiter is an error, not a silent lie.
    again = Workbook.wrap(raw)
    assert again._runner is wb._runner
    assert again.limiter is limiter
    assert Workbook.wrap(raw, limiter=limiter)._runner is wb._runner
    with pytest.raises(ValueError, match="different limiter"):
        Workbook.wrap(raw, limiter=anyio.CapacityLimiter(3))
    assert Workbook.wrap(openpyxl.Workbook()).limiter is None


def test_passing_a_raw_workbook_to_the_constructor_is_rejected() -> None:
    with pytest.raises(TypeError, match=r"Workbook\.wrap"):
        Workbook(openpyxl.Workbook())  # type: ignore[call-overload]  # ty: ignore[no-matching-overload]


def test_workbook_subclass_works() -> None:
    class Report(Workbook):
        __slots__ = ()

    report = Report()
    assert isinstance(report, Workbook)
    assert type(Report.wrap(openpyxl.Workbook())) is Report


# ------------------------------------------------- __setattr__ on subclasses
def test_setattr_subclass_class_attribute_counter_lives_on_the_wrapper() -> None:
    class Counting(Workbook):
        rows_written = 0

        def note_rows(self, n: int) -> None:
            self.rows_written += n

    wb = Counting()
    wb.note_rows(1)
    wb.note_rows(1)
    assert wb.rows_written == 2
    assert Counting.rows_written == 0  # the class attribute is untouched
    assert not hasattr(wb.wrapped, "rows_written")  # nothing leaked into openpyxl
    # A second instance starts from the class default again.
    assert Counting().rows_written == 0


def test_setattr_subclass_instance_attribute_stays_on_the_wrapper() -> None:
    class Annotated(Workbook):
        def __init__(self) -> None:
            super().__init__()
            self.note = "x"

    wb = Annotated()
    assert wb.note == "x"
    assert "note" in vars(wb)
    assert not hasattr(wb.wrapped, "note")
    wb.note = "y"
    assert wb.note == "y" and not hasattr(wb.wrapped, "note")
    # Subclasses of the sheet wrappers behave the same way.
    raw = openpyxl.Workbook()
    assert raw.active is not None

    class Tagged(AsyncWorksheet):
        def __init__(self, sheet: Any) -> None:
            super().__init__(sheet)
            self.tag = "t"

    ws = Tagged(raw.active)
    assert ws.tag == "t" and not hasattr(raw.active, "tag")


def test_setattr_existing_openpyxl_attributes_still_reach_the_raw_object() -> None:
    class Counting(Workbook):
        rows_written = 0

    wb = Counting()
    ws = wb.active
    assert ws is not None
    ws.title = "X"  # generated property: forwarded by its setter
    assert wb.wrapped.active is not None and wb.wrapped.active.title == "X"
    wb.iso_dates = True  # generated property on the workbook
    assert wb.wrapped.iso_dates is True
    # An attribute that exists on the raw object but has no generated property
    # goes to raw as well, even though the subclass could store it itself.
    # (Statically only the typed surface exists, hence setattr/getattr.)
    name = "custom_marker"
    setattr(wb.wrapped, name, 1)
    setattr(wb, name, 2)
    assert getattr(wb.wrapped, name) == 2
    assert name not in vars(wb)


def test_setattr_unknown_attribute_on_a_base_wrapper_lands_on_the_raw_object() -> None:
    # The base wrappers only define slots, so a new name has nowhere to live on
    # the wrapper and is put on the openpyxl object, as before.
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    name = "custom_flag"
    setattr(ws, name, True)
    assert getattr(ws.wrapped, name) is True
    assert getattr(ws, name) is True  # read back through ``__getattr__``
    setattr(wb, name, 3)
    assert getattr(wb.wrapped, name) == 3

    # A ``__slots__``-only subclass behaves like the base wrapper.
    class Slotted(Workbook):
        __slots__ = ()

    slotted = Slotted()
    setattr(slotted, name, 4)
    assert getattr(slotted.wrapped, name) == 4


async def test_deprecated_sheet_methods_warn_and_keep_working() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    with pytest.deprecated_call():
        assert wb.get_sheet_by_name("Sheet") is wb["Sheet"]  # same cached wrapper
    with pytest.deprecated_call():
        assert wb.get_sheet_names() == ["Sheet"]
    with pytest.deprecated_call():
        assert wb.get_index(ws) == 0
    with pytest.deprecated_call():
        wb.remove_sheet(ws)
    assert wb.sheetnames == []
    extra = wb.create_sheet("Extra")
    with pytest.deprecated_call():
        wb.create_named_range("rng", extra.wrapped, "A1")  # ty: ignore[deprecated]
    assert "rng" in wb.defined_names  # openpyxl stores it workbook-wide


# ---------------------------------------------------------------- error paths
async def test_missing_file_raises_file_not_found(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        await aioopenpyxl.load_workbook(tmp_path / "missing.xlsx")
    with pytest.raises(FileNotFoundError):
        async with aioopenpyxl.load_workbook(tmp_path / "missing.xlsx", read_only=True):
            pass  # pragma: no cover


async def test_corrupt_file_raises_bad_zip_file(tmp_path: Path) -> None:
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"this is not a zip archive")
    with pytest.raises(zipfile.BadZipFile):
        await aioopenpyxl.load_workbook(broken)
    with pytest.raises(zipfile.BadZipFile):
        await aioopenpyxl.load_workbook_bytes(b"this is not a zip archive")


async def test_load_workbook_result_cannot_be_awaited_twice(sample_xlsx: Path) -> None:
    pending = aioopenpyxl.load_workbook(sample_xlsx)
    wb = await pending
    try:
        with pytest.raises(RuntimeError):
            await pending
    finally:
        await wb.close()


async def test_to_bytes_on_read_only_workbook_raises(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        with pytest.raises(TypeError):
            await wb.to_bytes()


async def test_to_bytes_on_write_only_workbook() -> None:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("W")
    await ws.append_rows([[1, 2], [3, 4]])
    data = await wb.to_bytes()
    assert data[:2] == b"PK"
    async with aioopenpyxl.load_workbook_bytes(data, read_only=True) as loaded:
        assert await loaded["W"].read_rows() == [(1, 2), (3, 4)]


# ----------------------------------------------------------------- chartsheets
async def test_chartsheets_are_wrapped_and_share_the_runner() -> None:
    wb = aioopenpyxl.Workbook()
    cs = wb.create_chartsheet("Chart")
    assert isinstance(cs, AsyncChartsheet)
    assert isinstance(cs.wrapped, Chartsheet)
    cs.add_chart(BarChart())
    assert len(cast(Any, cs.wrapped)._charts) == 1  # private openpyxl state
    assert wb.chartsheets == [cs]
    assert cs.parent is wb
    assert wb.sheetnames.index("Chart") == 1  # openpyxl's index() only covers worksheets
    assert wb.sheetnames == ["Sheet", "Chart"]
    assert repr(cs) == "<AsyncChartsheet 'Chart'>"
    assert await cs.run(lambda raw: raw.title) == "Chart"


async def test_chartsheet_survives_a_save_and_reload_round_trip() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    await ws.append_rows([[1], [2], [3]])
    cs = wb.create_chartsheet("Chart")
    chart = BarChart()
    chart.add_data(Reference(ws.wrapped, min_col=1, min_row=1, max_row=3))
    cs.add_chart(chart)
    wb.active = cs
    data = await wb.to_bytes()

    async with aioopenpyxl.load_workbook_bytes(data) as loaded:
        assert loaded.sheetnames == ["Data", "Chart"]
        assert isinstance(loaded.chartsheets[0], AsyncChartsheet)
        assert loaded.chartsheets[0].title == "Chart"
        assert len(loaded.worksheets) == 1
        # The active sheet is the chartsheet; the wrapper type follows the raw type.
        active = loaded.active
        assert isinstance(active, AsyncChartsheet)
        assert active is loaded.chartsheets[0]
        assert active.parent is loaded


def test_chartsheet_from_tree_is_wrapped() -> None:
    node = fromstring(
        b'<chartsheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b"<sheetPr/><sheetViews><sheetView/></sheetViews></chartsheet>"
    )
    cs = AsyncChartsheet.from_tree(cast(Any, node))
    assert isinstance(cs, AsyncChartsheet)
    assert isinstance(cs.wrapped, Chartsheet)
    assert cs.parent is None
    assert cs.sheetViews.sheetView is not None
    assert AsyncChartsheet.from_tree(cast(Any, Element("chartsheet"))) is not None
