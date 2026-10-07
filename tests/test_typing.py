"""Static typing checks (run by ``ty``, nothing here executes at test time)."""

from __future__ import annotations

from collections.abc import AsyncGenerator
from typing import TYPE_CHECKING

import aioopenpyxl

if TYPE_CHECKING:
    from typing import Any

    import openpyxl
    from anyio import CapacityLimiter
    from openpyxl import _VisibilityType
    from openpyxl.cell import _CellGetValue, _CellOrMergedCell
    from openpyxl.cell.cell import Cell
    from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell
    from openpyxl.chartsheet.views import ChartsheetViewList
    from openpyxl.packaging.custom import CustomPropertyList
    from openpyxl.worksheet.page import PageMargins
    from typing_extensions import assert_type  # 3.10-compatible for the type checker

    # Read-only sheets yield ReadOnlyCell / EmptyCell, regular ones Cell / MergedCell.
    CellRow = tuple[_CellOrMergedCell | ReadOnlyCell | EmptyCell, ...]
    ValueRow = tuple[_CellGetValue, ...]


async def _overloads(ws: aioopenpyxl.AsyncWorksheet, values_only: bool) -> None:
    assert_type(ws.iter_rows(values_only=True), AsyncGenerator[ValueRow, None])
    assert_type(ws.iter_rows(), AsyncGenerator[CellRow, None])
    assert_type(ws.iter_rows(1, 2, 1, 1, False), AsyncGenerator[CellRow, None])
    # All five arguments positional, like openpyxl's own overloads.
    assert_type(ws.iter_rows(1, 2, 1, 2, True), AsyncGenerator[ValueRow, None])
    assert_type(ws.iter_rows(1, 2, 1, 2, True, chunk_size=8), AsyncGenerator[ValueRow, None])
    assert_type(ws.iter_rows(None, None, None, None, True), AsyncGenerator[ValueRow, None])
    assert_type(
        ws.iter_rows(values_only=values_only),
        AsyncGenerator[CellRow, None] | AsyncGenerator[ValueRow, None],
    )
    assert_type(ws.iter_cols(values_only=True, chunk_size=8), AsyncGenerator[ValueRow, None])
    assert_type(ws.iter_cols(prefetch=2), AsyncGenerator[CellRow, None])
    assert_type(ws.iter_cols(1, 2, 1, 1, False), AsyncGenerator[CellRow, None])
    assert_type(ws.iter_cols(1, 2, 1, 2, True), AsyncGenerator[ValueRow, None])
    assert_type(ws.iter_cols(1, 2, 1, 2, True, prefetch=1), AsyncGenerator[ValueRow, None])
    assert_type(
        ws.iter_cols(values_only=values_only),
        AsyncGenerator[CellRow, None] | AsyncGenerator[ValueRow, None],
    )
    assert_type(ws.values, AsyncGenerator[ValueRow, None])
    assert_type(ws.rows, AsyncGenerator[CellRow, None])
    assert_type(ws.columns, AsyncGenerator[CellRow, None])
    assert_type(ws.__aiter__(), AsyncGenerator[CellRow, None])
    assert_type(await ws.read_rows(), list[ValueRow])
    assert_type(await ws.read_rows(values_only=False), list[CellRow])
    assert_type(await ws.read_rows(1, 2, 1, 1, False), list[CellRow])
    assert_type(await ws.read_rows(1, 2, 1, 1, True), list[ValueRow])
    assert_type(await ws.read_rows(values_only=values_only), list[ValueRow] | list[CellRow])
    assert_type(ws.max_row, int | None)
    assert_type(ws.max_column, int | None)
    assert_type(ws.title, str)
    assert_type(ws.is_read_only, bool)
    assert_type(await ws.run(lambda raw: raw.title), str)
    assert_type(await ws.append_rows([[1, 2]]), int)
    assert_type(await ws.calculate_dimension(), str)
    assert_type(ws.cell(1, 1), _CellOrMergedCell)
    assert_type(ws.cell(1, 1, "x"), Cell)
    parent = ws.parent
    assert_type(parent, aioopenpyxl.Workbook | None)


async def _workbook(wb: aioopenpyxl.Workbook, raw: openpyxl.Workbook) -> None:
    assert_type(wb.active, aioopenpyxl.AsyncWorksheet | None)
    assert_type(wb["x"], aioopenpyxl.AsyncWorksheet)
    assert_type(wb.create_sheet("x"), aioopenpyxl.AsyncWorksheet)
    assert_type(wb.create_chartsheet(), aioopenpyxl.AsyncChartsheet)
    assert_type(wb.worksheets, list[aioopenpyxl.AsyncWorksheet])
    assert_type(wb.chartsheets, list[aioopenpyxl.AsyncChartsheet])
    assert_type(await wb.to_bytes(), bytes)
    assert_type(wb.read_only, bool)
    assert_type(wb.wrapped, openpyxl.Workbook)
    assert_type(wb.limiter, CapacityLimiter | None)
    assert_type(wb.custom_doc_props, CustomPropertyList[Any])
    assert_type(await wb.run(lambda raw: raw.sheetnames), list[str])
    assert_type(aioopenpyxl.Workbook(write_only=True), aioopenpyxl.Workbook)
    assert_type(aioopenpyxl.Workbook.wrap(raw), aioopenpyxl.Workbook)
    assert_type(aioopenpyxl.AsyncWorkbook.wrap(raw), aioopenpyxl.Workbook)
    assert_type(await aioopenpyxl.load_workbook("x.xlsx"), aioopenpyxl.Workbook)
    async with aioopenpyxl.load_workbook_bytes(b"") as loaded:
        assert_type(loaded, aioopenpyxl.Workbook)
    # Deprecated forwarders are still typed.
    assert_type(wb.get_sheet_by_name("x"), aioopenpyxl.AsyncWorksheet)


def _chartsheet(cs: aioopenpyxl.AsyncChartsheet, node: Any) -> None:
    # Descriptor-typed stub attributes are exposed as the value types they
    # produce, exactly as on the raw openpyxl class.
    assert_type(cs.pageMargins, PageMargins | None)
    cs.pageMargins = PageMargins(left=1.0)
    cs.pageMargins = None
    margins = cs.pageMargins
    if margins is not None:
        assert_type(margins.left, float)
    assert_type(cs.sheet_state, _VisibilityType)
    cs.sheet_state = "hidden"
    assert_type(cs.sheetViews, ChartsheetViewList)  # Typed[..., Literal[False]]
    assert_type(cs.title, str)
    assert_type(aioopenpyxl.AsyncChartsheet.from_tree(node), aioopenpyxl.AsyncChartsheet | None)


def _module_constants() -> None:
    # Explicit re-exports: resolvable without the runtime-only ``__getattr__``.
    assert_type(aioopenpyxl.LXML, bool)
    assert_type(aioopenpyxl.DEBUG, bool)
    assert_type(aioopenpyxl.__author__, str)


def _shims() -> None:
    # The generated ``.pyi`` shims resolve the aliased sub-packages statically.
    from aioopenpyxl.styles import Font
    from aioopenpyxl.utils import get_column_letter

    assert_type(get_column_letter(1), str)
    assert_type(Font(), Font)


def test_module_is_importable() -> None:
    assert aioopenpyxl.AsyncWorksheet is not None
