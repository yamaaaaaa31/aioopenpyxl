"""Real spawned workers: Excel round trips, admission, cancellation and shutdown."""

from __future__ import annotations

import asyncio
import io
import multiprocessing
import os
import pickle
import time
from concurrent.futures.process import BrokenProcessPool
from pathlib import Path
from typing import Any

import anyio
import anyio.lowlevel
import openpyxl
import pytest

from aioopenpyxl import ProcessPool

pytestmark = pytest.mark.anyio


# Module-level functions are importable by spawn workers on every platform.
def _export(rows: list[list[str | int]], *, title: str) -> tuple[int, bytes]:
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(title)
    for row in rows:
        ws.append(row)
    output = io.BytesIO()
    wb.save(output)
    return os.getpid(), output.getvalue()


def _read(source: Path | bytes) -> list[tuple[Any, ...]]:
    wb = openpyxl.load_workbook(
        io.BytesIO(source) if isinstance(source, bytes) else source, read_only=True
    )
    try:
        return list(wb.worksheets[0].values)
    finally:
        wb.close()


def _fail() -> None:
    raise ValueError("invalid Excel job")


def _mark(path: Path) -> int:
    path.write_text("ran")
    return os.getpid()


def _busy_job(started: Path, release: Path, finished: Path) -> int:
    # Publish the PID atomically: exists() must not observe an empty file
    # between opening the signal file and writing its contents.
    pending = started.with_suffix(".pending")
    pending.write_text(str(os.getpid()))
    pending.replace(started)
    deadline = time.monotonic() + 10
    while not release.exists():
        # Real Python CPU work in the child, not a sleeping stand-in. The
        # parent must keep running to create the release file.
        sum(i * i for i in range(20_000))
        if time.monotonic() > deadline:
            raise TimeoutError("parent did not release the process job")
    finished.write_text("finished")
    return os.getpid()


async def _wait_for_file(path: Path) -> None:
    with anyio.fail_after(8):
        # The child signals via a local temporary file, not an anyio Event.
        while not path.exists():  # noqa: ASYNC110, ASYNC240
            await anyio.sleep(0.005)


def _assert_joined(*pids: int) -> None:
    assert not set(pids).intersection(p.pid for p in multiprocessing.active_children())


async def test_excel_jobs_run_in_a_reused_child_and_round_trip(tmp_path: Path) -> None:
    async with ProcessPool() as pool:
        pid, data = await pool.run(_export, [["name", "count"], ["item", 7]], title="Data")
        assert pid != os.getpid()
        assert await pool.run(os.getpid) == pid
        assert await pool.run(_read, data) == [("name", "count"), ("item", 7)]
        path = tmp_path / "export.xlsx"
        path.write_bytes(data)
        assert await pool.run(_read, path) == [("name", "count"), ("item", 7)]
    _assert_joined(pid)


async def test_worker_and_serialization_errors_propagate_without_breaking_pool() -> None:
    async with ProcessPool() as pool:
        with pytest.raises(ValueError, match="invalid Excel job"):
            await pool.run(_fail)
        with pytest.raises((AttributeError, pickle.PicklingError)):
            await pool.run(lambda: None)
        assert await pool.run(os.getpid) != os.getpid()


async def test_two_cpu_jobs_run_in_distinct_processes_while_loop_progresses(tmp_path: Path) -> None:
    started = [tmp_path / f"start{i}" for i in range(2)]
    finished = [tmp_path / f"done{i}" for i in range(2)]
    release = tmp_path / "release"
    async with ProcessPool(max_workers=2) as pool, anyio.create_task_group() as tg:
        for i in range(2):
            tg.start_soon(pool.run, _busy_job, started[i], release, finished[i])
        try:
            for path in started:
                await _wait_for_file(path)
            pids = [int(path.read_text()) for path in started]
            assert len({*pids, os.getpid()}) == 3
            assert all(not path.exists() for path in finished)
        finally:
            release.touch()
    assert all(path.exists() for path in finished)
    _assert_joined(*pids)


async def test_cancelling_a_queued_job_does_not_submit_it(tmp_path: Path) -> None:
    started, release, finished, queued = [
        tmp_path / name for name in ("started", "release", "finished", "queued")
    ]
    scope = anyio.CancelScope()

    async def queued_job(pool: ProcessPool) -> None:
        with scope:
            await pool.run(_mark, queued)

    async with ProcessPool() as pool, anyio.create_task_group() as tg:
        tg.start_soon(pool.run, _busy_job, started, release, finished)
        try:
            await _wait_for_file(started)
            tg.start_soon(queued_job, pool)
            with anyio.fail_after(5):
                # Observe a waiter owned by the pool, not a test-owned Event.
                while not pool._limiter.statistics().tasks_waiting:  # noqa: ASYNC110
                    await anyio.sleep(0.005)
            scope.cancel()
        finally:
            release.touch()
    assert scope.cancelled_caught
    assert finished.exists()
    assert not queued.exists()


async def test_anyio_cancellation_waits_for_job_and_keeps_pool_usable(tmp_path: Path) -> None:
    started, release, finished = [tmp_path / name for name in ("started", "release", "finished")]
    scope = anyio.CancelScope()
    returned = anyio.Event()

    async def job(pool: ProcessPool) -> None:
        with scope:
            await pool.run(_busy_job, started, release, finished)
            await anyio.lowlevel.checkpoint()
        returned.set()

    async with ProcessPool() as pool:
        async with anyio.create_task_group() as tg:
            tg.start_soon(job, pool)
            try:
                await _wait_for_file(started)
                scope.cancel()
                await anyio.sleep(0.02)
                assert not returned.is_set()
            finally:
                release.touch()
        assert scope.cancelled_caught
        assert finished.exists()
        assert await pool.run(os.getpid) == int(started.read_text())


@pytest.mark.parametrize("anyio_backend", ["asyncio"])
async def test_native_cancellation_waits_for_job_even_when_repeated(tmp_path: Path) -> None:
    started, release, finished = [tmp_path / name for name in ("started", "release", "finished")]
    async with ProcessPool() as pool:
        task = asyncio.create_task(pool.run(_busy_job, started, release, finished))
        try:
            await _wait_for_file(started)
            task.cancel()
            await anyio.sleep(0.02)
            assert not task.done()
            task.cancel()
        finally:
            release.touch()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert finished.exists()
        assert await pool.run(os.getpid) == int(started.read_text())


async def test_concurrent_close_waits_for_jobs_and_rejects_new_work(tmp_path: Path) -> None:
    started, release, finished = [tmp_path / name for name in ("started", "release", "finished")]
    closing = anyio.Event()
    closed = anyio.Event()

    async def close(pool: ProcessPool) -> None:
        closing.set()
        await pool.aclose()
        closed.set()

    async with ProcessPool() as pool, anyio.create_task_group() as tg:
        tg.start_soon(pool.run, _busy_job, started, release, finished)
        try:
            await _wait_for_file(started)
            tg.start_soon(close, pool)
            await closing.wait()
            tg.start_soon(pool.aclose)
            with pytest.raises(RuntimeError, match="closed"):
                await pool.run(os.getpid)
            assert not closed.is_set()
        finally:
            release.touch()
    assert finished.exists()
    _assert_joined(int(started.read_text()))


async def test_cancelled_context_exit_still_joins_workers() -> None:
    with anyio.CancelScope() as scope:
        async with ProcessPool() as pool:
            pid = await pool.run(os.getpid)
            scope.cancel()
    _assert_joined(pid)


@pytest.mark.parametrize("anyio_backend", ["asyncio"])
async def test_native_cancellation_during_close_still_joins_workers(tmp_path: Path) -> None:
    started, release, finished = [tmp_path / name for name in ("started", "release", "finished")]
    async with ProcessPool() as pool:
        job = asyncio.create_task(pool.run(_busy_job, started, release, finished))
        try:
            await _wait_for_file(started)
            close = asyncio.create_task(pool.aclose())
            await anyio.sleep(0.02)
            close.cancel()
            await anyio.sleep(0.02)
            assert not close.done()
            close.cancel()
        finally:
            release.touch()
        with pytest.raises(asyncio.CancelledError):
            await close
        await job
        assert finished.exists()
        _assert_joined(int(started.read_text()))


async def test_worker_crash_is_reported_and_jobs_are_not_retried() -> None:
    async with ProcessPool() as pool:
        with pytest.raises(BrokenProcessPool):
            await pool.run(os._exit, 17)
        with pytest.raises(BrokenProcessPool):
            await pool.run(os.getpid)


async def test_pool_lifecycle_is_explicit_and_close_is_idempotent() -> None:
    pool = ProcessPool()
    with pytest.raises(RuntimeError, match="async with"):
        await pool.run(os.getpid)
    async with pool:
        with pytest.raises(RuntimeError, match="already open"):
            await pool.__aenter__()
    await pool.aclose()
    with pytest.raises(RuntimeError, match="closed"):
        await pool.__aenter__()
    with pytest.raises(RuntimeError, match="closed"):
        await pool.run(os.getpid)
    unopened = ProcessPool()
    await unopened.aclose()
    with pytest.raises(RuntimeError, match="closed"):
        await unopened.__aenter__()


@pytest.mark.parametrize("count", [0, -1])
def test_invalid_worker_count(count: int) -> None:
    with pytest.raises(ValueError, match="max_workers"):
        ProcessPool(max_workers=count)
