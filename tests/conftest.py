from __future__ import annotations

import gc
import logging
import threading
from pathlib import Path

import anyio
import openpyxl
import pytest

import aioopenpyxl

ROWS: list[tuple[str | int | float, ...]] = [
    ("id", "name", "score"),
    *[(i, f"user{i}", i * 1.5) for i in range(1, 51)],
]

# trio warns when an async generator is garbage collected before being exhausted
# (asyncio does not).  Tests that abandon generators on purpose carry this mark.
abandons_generator = pytest.mark.filterwarnings("ignore:Async generator:ResourceWarning")


def _openpyxl_version() -> tuple[int, ...]:
    return tuple(int(part) for part in openpyxl.__version__.split(".")[:3] if part.isdigit())


def pytest_configure(config: pytest.Config) -> None:
    if _openpyxl_version() < (3, 1, 3):
        # openpyxl <= 3.1.2 does not close the zip member it reads a read-only
        # sheet from when the row generator is closed early (``ws.fetch(key)``,
        # breaking out of ``iter_rows`` ...), so the archive's file handle is only
        # released by the garbage collector -- with a ResourceWarning that
        # ``filterwarnings = error`` would turn into a random test failure.
        # Fixed upstream in 3.1.3 (``with self._get_source() as src``).
        config.addinivalue_line(
            "filterwarnings",
            r"ignore:unclosed file <_io\.BufferedReader name='.*\.xlsx'>:ResourceWarning",
        )


@pytest.fixture(params=["asyncio", "trio"])
def anyio_backend(request: pytest.FixtureRequest) -> str:
    """Run every async test on both backends."""
    return request.param


@pytest.fixture
def sample_xlsx(tmp_path: Path) -> Path:
    """A plain workbook written with raw openpyxl: 1 header + 50 data rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    for row in ROWS:
        ws.append(row)
    wb.create_sheet("Empty")
    path = tmp_path / "sample.xlsx"
    wb.save(path)
    return path


# ------------------------------------------------------------------ helpers
def prefetch_threads() -> list[threading.Thread]:
    """The ``prefetch`` producer threads that are currently alive."""
    return [t for t in threading.enumerate() if t.name == "aioopenpyxl-prefetch"]


async def wait_until_idle(wb: aioopenpyxl.Workbook) -> None:
    """Give the event loop's async-generator finaliser time to run ``aclose()``."""
    with anyio.fail_after(5):
        while wb._runner.busy or prefetch_threads():
            gc.collect()
            await anyio.sleep(0.01)


def no_errors_logged(caplog: pytest.LogCaptureFixture) -> bool:
    # trio logs "Exception ignored during finalization of async generator" at
    # ERROR level; asyncio hands task exceptions to anyio, which re-raises them.
    return not [r for r in caplog.records if r.levelno >= logging.ERROR]
