"""The ``Runner``: thread hops, serialisation, the busy guard, the registry and limiters."""

from __future__ import annotations

import contextlib
import contextvars
import gc
import threading
import time
from pathlib import Path
from typing import Any, cast

import anyio
import anyio.to_thread
import openpyxl
import pytest
from openpyxl.chartsheet import Chartsheet

import aioopenpyxl
import aioopenpyxl._executor
from aioopenpyxl import (
    AsyncChartsheet,
    AsyncWorksheet,
    Runner,
    Workbook,
    WorkbookBusyError,
    run_sync,
)
from aioopenpyxl._executor import runner_for
from tests.conftest import prefetch_threads

pytestmark = pytest.mark.anyio

var: contextvars.ContextVar[str] = contextvars.ContextVar("var", default="unset")


class _Concurrency:
    """Counts how many ``work()`` calls overlap."""

    def __init__(self) -> None:
        self.active = 0
        self.max_active = 0
        self._lock = threading.Lock()

    def work(self) -> None:
        with self._lock:
            self.active += 1
            self.max_active = max(self.max_active, self.active)
        time.sleep(0.02)
        with self._lock:
            self.active -= 1


# ------------------------------------------------------------------- run_sync
async def test_run_sync_runs_off_the_event_loop_thread() -> None:
    main = threading.get_ident()
    worker = await run_sync(None, threading.get_ident)
    assert worker != main


async def test_run_sync_propagates_contextvars() -> None:
    var.set("hello")
    assert await run_sync(None, var.get) == "hello"


async def test_run_sync_propagates_exceptions() -> None:
    def boom() -> None:
        raise ValueError("boom")

    with pytest.raises(ValueError, match="boom"):
        await run_sync(None, boom)


async def test_run_sync_respects_the_capacity_limiter() -> None:
    limiter = anyio.CapacityLimiter(1)
    probe = _Concurrency()
    async with anyio.create_task_group() as tg:
        for _ in range(4):
            tg.start_soon(run_sync, limiter, probe.work)
    assert probe.max_active == 1


async def test_without_limiter_threads_run_in_parallel() -> None:
    # All four threads must be alive at once for the barrier to open; if the
    # calls were serialised, ``wait`` would time out with BrokenBarrierError.
    barrier = threading.Barrier(4, timeout=5)
    async with anyio.create_task_group() as tg:
        for _ in range(4):
            tg.start_soon(run_sync, None, barrier.wait)
    assert barrier.broken is False


# --------------------------------------------------------------------- Runner
async def test_runner_serialises_calls() -> None:
    runner = Runner()
    probe = _Concurrency()
    async with anyio.create_task_group() as tg:
        for _ in range(5):
            tg.start_soon(runner, probe.work)
    assert probe.max_active == 1
    assert runner.busy is False


def test_runner_busy_is_readable_outside_an_event_loop() -> None:
    # The semaphore is created lazily, so this must not raise NoEventLoopError.
    assert Runner().busy is False
    assert Runner(anyio.CapacityLimiter(1)).busy is False


async def test_runner_as_context_manager_marks_the_workbook_busy() -> None:
    runner = Runner()
    assert runner.busy is False
    async with runner:
        assert runner.busy is True
        # Thread hops are still possible for the holder.
        assert await runner.offload(threading.get_ident) != threading.get_ident()
    assert runner.busy is False


async def test_runner_can_be_released_from_another_task() -> None:
    # The finaliser of an abandoned ``async for`` runs in a different task than
    # the one that acquired the runner; anyio.Lock would raise here.
    runner = Runner()
    await runner.acquire()
    assert runner.busy is True

    async def release_from_elsewhere() -> None:
        runner.release()

    async with anyio.create_task_group() as tg:
        tg.start_soon(release_from_elsewhere)
    assert runner.busy is False
    # And the runner is fully usable afterwards.
    assert await runner(lambda: 42) == 42


async def test_runner_release_without_acquire_is_an_error() -> None:
    with pytest.raises(RuntimeError, match="acquire"):
        Runner().release()


async def test_runner_waits_for_the_holder_before_running() -> None:
    runner = Runner()
    order: list[str] = []

    async def hold() -> None:
        async with runner:
            order.append("hold-start")
            await anyio.sleep(0.05)
            order.append("hold-end")

    async with anyio.create_task_group() as tg:
        tg.start_soon(hold)
        await anyio.sleep(0.01)
        await runner(order.append, "call")
    assert order == ["hold-start", "hold-end", "call"]


# ----------------------------------------------------------------- busy guard
async def test_sync_access_while_a_thread_owns_the_workbook_raises() -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    started = threading.Event()
    release = threading.Event()

    def slow(raw: openpyxl.Workbook) -> None:
        started.set()
        assert release.wait(5)

    async with anyio.create_task_group() as tg:
        tg.start_soon(wb.run, slow)
        await anyio.to_thread.run_sync(started.wait, 5)

        with pytest.raises(WorkbookBusyError):
            ws.cell(1, 1, 1)
        with pytest.raises(WorkbookBusyError):
            ws["A1"] = 1
        with pytest.raises(WorkbookBusyError):
            ws.title = "x"
        with pytest.raises(WorkbookBusyError):
            _ = wb.sheetnames
        with pytest.raises(WorkbookBusyError):
            _ = wb.active
        assert ws.wrapped is wb.wrapped.active  # the escape hatch is never guarded

        release.set()
    ws.cell(1, 1, 1)  # fine again once the thread is done
    assert ws["A1"].value == 1


async def test_repr_does_not_bypass_the_busy_guard() -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    cs = wb.create_chartsheet("C")
    started = threading.Event()
    release = threading.Event()

    def slow(r: openpyxl.Workbook) -> None:
        started.set()
        release.wait(5)

    async with anyio.create_task_group() as tg:
        tg.start_soon(wb.run, slow)
        await anyio.to_thread.run_sync(started.wait, 5)
        assert repr(wb) == "<Workbook busy>"
        assert repr(ws) == "<AsyncWorksheet busy>"
        assert repr(cs) == "<AsyncChartsheet busy>"
        release.set()
    assert repr(wb) == "<Workbook sheets=['Sheet', 'C']>"
    assert repr(ws) == "<AsyncWorksheet 'Sheet'>"
    assert repr(cs) == "<AsyncChartsheet 'C'>"


# ------------------------------------------------------------- shared runner
async def test_wrapping_the_same_workbook_twice_shares_the_runner() -> None:
    wb = aioopenpyxl.Workbook()
    raw = wb.wrapped
    started = threading.Event()
    release = threading.Event()

    def slow(r: openpyxl.Workbook) -> None:
        started.set()
        release.wait(5)

    async with anyio.create_task_group() as tg:
        tg.start_soon(wb.run, slow)
        await anyio.to_thread.run_sync(started.wait, 5)
        other = Workbook.wrap(raw)
        assert other._runner is wb._runner
        with pytest.raises(WorkbookBusyError):
            cast(Any, other).active.cell(1, 1, 1)
        loose = AsyncWorksheet(cast(Any, raw).active)
        assert loose._runner is wb._runner
        with pytest.raises(WorkbookBusyError):
            loose.cell(1, 1, 1)
        assert repr(loose) == "<AsyncWorksheet busy>"
        release.set()
    assert runner_for(raw) is wb._runner
    assert cast(Any, Workbook.wrap(raw)).active.cell(1, 1, 1).value == 1


def test_explicit_runner_goes_through_the_registry() -> None:
    raw = openpyxl.Workbook()
    runner = Runner()
    # First wrapper with an explicit runner registers it for the workbook ...
    loose = AsyncWorksheet(cast(Any, raw).active, runner=runner)
    assert loose._runner is runner
    assert Workbook.wrap(raw)._runner is runner
    assert AsyncWorksheet(cast(Any, raw).active)._runner is runner
    assert AsyncChartsheet(raw.create_chartsheet("C"), runner=runner)._runner is runner
    # ... and a *different* explicit runner for the same workbook is refused.
    with pytest.raises(ValueError, match="different Runner"):
        AsyncWorksheet(cast(Any, raw).active, runner=Runner())
    with pytest.raises(ValueError, match="different Runner"):
        AsyncChartsheet(raw.chartsheets[0], runner=Runner())
    with pytest.raises(ValueError, match="different Runner"):
        Workbook._from_runner(raw, Runner())
    # Parentless sheets keep whatever runner they are given.
    assert AsyncChartsheet(Chartsheet(), runner=runner)._runner is runner


async def test_explicit_runner_wrapper_is_guarded_by_the_shared_lock() -> None:
    raw = openpyxl.Workbook()
    loose = AsyncWorksheet(cast(Any, raw).active, runner=Runner())
    wb = Workbook.wrap(raw)
    started = threading.Event()
    release = threading.Event()

    def slow(r: openpyxl.Workbook) -> None:
        started.set()
        release.wait(5)

    async with anyio.create_task_group() as tg:
        tg.start_soon(wb.run, slow)
        await anyio.to_thread.run_sync(started.wait, 5)
        with pytest.raises(WorkbookBusyError):
            loose.cell(1, 1, 1)
        release.set()
    assert loose.cell(1, 1, 1).value == 1


def test_runner_registry_is_weak() -> None:
    raw = openpyxl.Workbook()
    runner = Workbook.wrap(raw)._runner
    assert runner_for(raw) is runner
    del raw
    gc.collect()
    # No way to look the entry up any more; a new workbook gets a fresh runner.
    assert runner_for(openpyxl.Workbook()) is not runner


def test_loose_chartsheet_shares_its_workbook_runner() -> None:
    wb = aioopenpyxl.Workbook()
    cs = wb.create_chartsheet("C")
    assert AsyncChartsheet(cs.wrapped)._runner is wb._runner
    assert AsyncChartsheet(cs.wrapped).parent is not None
    # Parentless sheets get their own runner.
    orphan = AsyncChartsheet(Chartsheet())
    assert orphan._runner is not wb._runner


# -------------------------------------------------------------------- limiter
async def test_limiter_is_passed_to_every_thread_hop(
    sample_xlsx: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Every hop borrows a token of the configured limiter on behalf of the
    # call (the Runner manages the token itself so that a cancellation cannot
    # hand it back while the thread is still running).
    seen: list[anyio.CapacityLimiter] = []
    real = aioopenpyxl._executor._acquire_token

    async def spy(limiter: anyio.CapacityLimiter, borrower: object, **kw: Any) -> Any:
        seen.append(limiter)
        return await real(limiter, borrower, **kw)

    monkeypatch.setattr(aioopenpyxl._executor, "_acquire_token", spy)

    limiter = anyio.CapacityLimiter(1)
    async with aioopenpyxl.load_workbook(sample_xlsx, limiter=limiter) as wb:
        await wb.run(lambda raw: None)
        await wb["Data"].read_rows(max_row=1)
        await wb.to_bytes()
    # load, run, read_rows, to_bytes, close
    assert len(seen) == 5 and all(lim is limiter for lim in seen)
    assert limiter.borrowed_tokens == 0

    seen.clear()
    wb2 = Workbook(limiter=limiter)
    await wb2.to_bytes()
    assert seen == [limiter]

    seen.clear()
    await Workbook().to_bytes()
    assert seen == [anyio.to_thread.current_default_thread_limiter()]


class _Parsers:
    """Counts how many wrapped ``iter_rows`` generators are alive at once."""

    def __init__(self) -> None:
        self.active = 0
        self.peak = 0
        self.lock = threading.Lock()

    def wrap(self, ws: AsyncWorksheet) -> None:
        raw: Any = ws.wrapped
        real = raw.iter_rows

        def counted(**kwargs: Any) -> Any:
            with self.lock:
                self.active += 1
                self.peak = max(self.peak, self.active)
            try:
                yield from real(**kwargs)
            finally:
                with self.lock:
                    self.active -= 1

        raw.iter_rows = counted


@pytest.mark.parametrize("shared", [True, False])
async def test_prefetch_producers_respect_a_shared_limiter(shared: bool) -> None:
    limiter = anyio.CapacityLimiter(1) if shared else None
    parsers = _Parsers()
    sheets: list[AsyncWorksheet] = []
    for _ in range(4):
        wb = aioopenpyxl.Workbook(limiter=limiter)
        ws = wb.active
        assert ws is not None
        await ws.append_rows([[i] for i in range(40)])
        parsers.wrap(ws)
        sheets.append(ws)
    results: list[list[Any]] = []

    async def consume(ws: AsyncWorksheet) -> None:
        rows = []
        async for row in ws.iter_rows(values_only=True, chunk_size=5, prefetch=2):
            rows.append(row)
            await anyio.sleep(0.001)
        results.append(rows)

    async with anyio.create_task_group() as tg:
        for sheet in sheets:
            tg.start_soon(consume, sheet)
    assert results == [[(i,) for i in range(40)]] * 4
    if shared:
        # One token: one producer parsing at a time, even across four workbooks.
        assert parsers.peak == 1
        assert limiter is not None and limiter.borrowed_tokens == 0
    else:
        assert parsers.peak >= 1
    for sheet in sheets:
        assert sheet._runner.busy is False
    assert not prefetch_threads()


async def test_prefetch_token_is_returned_when_the_loop_is_left_early() -> None:
    limiter = anyio.CapacityLimiter(1)
    wb = aioopenpyxl.Workbook(limiter=limiter)
    ws = wb.active
    assert ws is not None
    await ws.append_rows([[i] for i in range(40)])
    rows = ws.iter_rows(chunk_size=3, prefetch=2)
    await anext(rows)
    assert limiter.borrowed_tokens == 1  # held on the producer's behalf
    await rows.aclose()
    assert limiter.borrowed_tokens == 0
    assert wb._runner.busy is False
    # A cancelled loop (closed synchronously through ``aclosing``) returns it too.
    with anyio.move_on_after(0.01):
        async with contextlib.aclosing(ws.iter_rows(chunk_size=3, prefetch=2)) as cancelled:
            async for _ in cancelled:
                await anyio.sleep(10)
    assert limiter.borrowed_tokens == 0
    assert await ws.read_rows(max_row=1) == [(0,)]
