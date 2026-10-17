"""Opt-in persistent process workbooks through the existing async entry points."""

from __future__ import annotations

import asyncio
import io
import multiprocessing
import os
import threading
from concurrent.futures.process import BrokenProcessPool
from pathlib import Path
from typing import Any

import anyio
import openpyxl
import pytest
from openpyxl.styles import Font

import aioopenpyxl
from tests.test_process import _assert_joined, _busy_job, _wait_for_file

pytestmark = pytest.mark.anyio


async def test_constructor_keeps_state_in_one_child_and_saves(tmp_path: Path) -> None:
    wb = aioopenpyxl.Workbook(process=True)
    assert isinstance(wb, aioopenpyxl.Workbook)
    assert repr(wb) == "<Workbook process>"

    def fill(raw: openpyxl.Workbook) -> int:
        ws = raw.active
        assert ws is not None
        ws.title = "Report"
        ws["A1"] = "header"
        ws["A1"].font = Font(bold=True)
        ws["B1"] = "=1+2"
        return os.getpid()

    async with wb:
        pid = await wb.run(fill)
        assert pid != os.getpid()
        assert await wb.run(lambda raw: (os.getpid(), raw.sheetnames)) == (pid, ["Report"])
        output = tmp_path / "report.xlsx"
        await wb.save(output)
        data = await wb.to_bytes()
        stream = io.BytesIO()
        await wb.save(stream)
    _assert_joined(pid)
    for source in (output, io.BytesIO(data), stream):
        async with aioopenpyxl.load_workbook(source) as restored:
            ws = restored["Report"]
            assert ws["A1"].value == "header"
            assert ws["A1"].font.bold is True
            assert ws["B1"].value == "=1+2"


@pytest.mark.parametrize("read_only", [False, True])
async def test_load_and_bytes_inputs_support_read_only_and_local_callbacks(
    sample_xlsx: Path, read_only: bool
) -> None:
    data = await anyio.Path(sample_xlsx).read_bytes()
    sources = (sample_xlsx, io.BytesIO(data))
    for source in sources:
        if isinstance(source, io.BytesIO):
            source.seek(7)
        async with aioopenpyxl.load_workbook(source, read_only=read_only, process=True) as wb:
            pid = await wb.run(lambda raw: os.getpid())
            assert pid != os.getpid()
            assert await wb.run(lambda raw: raw.read_only) == read_only
            assert await wb.run(lambda raw: list(raw["Data"].values)[:2]) == [
                ("id", "name", "score"),
                (1, "user1", 1.5),
            ]
        _assert_joined(pid)
        if isinstance(source, io.BytesIO):
            assert source.tell() == 7
    wb = await aioopenpyxl.load_workbook_bytes(data, process=True)
    try:
        assert await wb.run(lambda raw: raw.sheetnames) == ["Data", "Empty"]
    finally:
        await wb.close()


async def test_write_only_workbook_is_created_and_filled_in_child() -> None:
    def append(raw: openpyxl.Workbook) -> None:
        raw.create_sheet("Data").append(["item", 42])

    async with aioopenpyxl.Workbook(write_only=True, process=True) as wb:
        await wb.run(append)
        data = await wb.to_bytes()
    async with aioopenpyxl.load_workbook_bytes(data, read_only=True) as restored:
        assert await restored["Data"].read_rows() == [("item", 42)]


async def test_closing_unsaved_write_only_workbook_removes_spool_file() -> None:
    def append(raw: openpyxl.Workbook) -> str:
        ws = raw.create_sheet("Data")
        ws.append(["item", 42])
        return ws._writer.out

    async with aioopenpyxl.Workbook(write_only=True, process=True) as wb:
        path = await wb.run(append)
        assert await anyio.Path(path).exists()
    assert not await anyio.Path(path).exists()


async def test_sync_access_is_rejected_without_process_messages() -> None:
    async with aioopenpyxl.Workbook(process=True) as wb:
        for read in (lambda: wb.active, lambda: wb.sheetnames, lambda: wb["Sheet"]):
            with pytest.raises(aioopenpyxl.BlockingCallError, match=r"wb\.run"):
                read()
        with pytest.raises(aioopenpyxl.BlockingCallError, match=r"wb\.run"):
            _ = wb.wrapped
        with pytest.raises(aioopenpyxl.BlockingCallError, match=r"wb\.run"):
            wb.iso_dates = True
        with pytest.raises(aioopenpyxl.BlockingCallError, match=r"wb\.run"):
            wb.create_sheet("Other")
        assert await wb.run(lambda raw: raw.sheetnames) == ["Sheet"]


async def test_independent_process_workbooks_do_not_share_state() -> None:
    async with (
        aioopenpyxl.Workbook(process=True) as first,
        aioopenpyxl.Workbook(process=True) as second,
    ):
        await first.run(lambda raw: raw.create_sheet("First").title)
        assert await first.run(lambda raw: raw.sheetnames) == ["Sheet", "First"]
        assert await second.run(lambda raw: raw.sheetnames) == ["Sheet"]
        pids = [await book.run(lambda raw: os.getpid()) for book in (first, second)]
        assert pids[0] != pids[1]
    _assert_joined(*pids)


async def test_callback_errors_and_unserializable_closures_leave_workbook_usable() -> None:
    class InvalidCell(ValueError):
        pass

    def fail(raw: openpyxl.Workbook) -> None:
        raise InvalidCell("bad value")

    async with aioopenpyxl.Workbook(process=True) as wb:
        with pytest.raises(InvalidCell, match="bad value"):
            await wb.run(fail)
        lock = threading.Lock()
        with pytest.raises(TypeError, match="pickle"):
            await wb.run(lambda raw: lock.locked())
        assert await wb.run(lambda raw: raw.sheetnames) == ["Sheet"]


async def test_async_callback_is_rejected() -> None:
    async def callback(raw: openpyxl.Workbook) -> None:
        pass

    async with aioopenpyxl.Workbook(process=True) as wb:
        with pytest.raises(TypeError, match="synchronous"):
            _ = await wb.run(callback)


async def test_invalid_load_cleans_up_its_worker(tmp_path: Path) -> None:
    before = {p.pid for p in multiprocessing.active_children()}
    with pytest.raises(FileNotFoundError):
        await aioopenpyxl.load_workbook(tmp_path / "missing.xlsx", process=True)
    assert {p.pid for p in multiprocessing.active_children()} == before


async def test_close_is_idempotent_and_does_not_start_an_unused_worker() -> None:
    wb = aioopenpyxl.Workbook(process=True)
    backend = wb._process_backend
    assert backend is not None and backend._executor is None
    await wb.close()
    await wb.close()
    assert backend._executor is None
    with pytest.raises(RuntimeError, match="closed"):
        await wb.run(lambda raw: raw.sheetnames)


async def test_relative_save_path_is_resolved_in_the_parent(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    child_directory = tmp_path / "child"
    child_directory.mkdir()
    async with aioopenpyxl.Workbook(process=True) as wb:
        # A callback may change its child's cwd without changing the caller's.
        await wb.run(lambda raw: os.chdir(child_directory))
        await wb.save("report.xlsx")
    assert await anyio.Path(tmp_path / "report.xlsx").exists()
    assert not await anyio.Path(child_directory / "report.xlsx").exists()


async def test_worker_crash_does_not_restart_workbook_or_mask_error_on_close() -> None:
    async with aioopenpyxl.Workbook(process=True) as wb:
        pid = await wb.run(lambda raw: os.getpid())
        with pytest.raises(BrokenProcessPool):
            await wb.run(lambda raw: os._exit(17))
        with pytest.raises(BrokenProcessPool):
            await wb.run(lambda raw: raw.sheetnames)
    _assert_joined(pid)


async def test_anyio_cancelled_exit_joins_worker() -> None:
    with anyio.CancelScope() as scope:
        async with aioopenpyxl.Workbook(process=True) as wb:
            pid = await wb.run(lambda raw: os.getpid())
            scope.cancel()
    _assert_joined(pid)


@pytest.mark.parametrize("anyio_backend", ["asyncio"])
async def test_native_cancel_waits_for_cpu_work_and_preserves_workbook(tmp_path: Path) -> None:
    started, release, finished = [tmp_path / name for name in ("started", "release", "finished")]
    async with aioopenpyxl.Workbook(process=True) as wb:
        task = asyncio.create_task(wb.run(lambda raw: _busy_job(started, release, finished)))
        try:
            await _wait_for_file(started)
            task.cancel()
            await anyio.sleep(0.02)
            assert not task.done()
            task.cancel()
        finally:
            release.touch()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert finished.exists()
        assert await wb.run(lambda raw: os.getpid()) == int(started.read_text())
    _assert_joined(int(started.read_text()))


@pytest.mark.parametrize("anyio_backend", ["asyncio"])
async def test_cancelled_load_waits_and_joins_the_unreturned_workbook(sample_xlsx: Path) -> None:
    started = threading.Event()
    release = threading.Event()

    class SlowInput(io.BytesIO):
        def read(self, *args: Any, **kwargs: Any) -> bytes:
            started.set()
            assert release.wait(8)
            return super().read(*args, **kwargs)

    before = {p.pid for p in multiprocessing.active_children()}
    data = await anyio.Path(sample_xlsx).read_bytes()

    async def load() -> None:
        await aioopenpyxl.load_workbook(SlowInput(data), process=True)

    task = asyncio.create_task(load())
    try:
        assert await anyio.to_thread.run_sync(started.wait, 5)
        task.cancel()
        await anyio.sleep(0.02)
        assert not task.done()
    finally:
        release.set()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert {p.pid for p in multiprocessing.active_children()} == before
