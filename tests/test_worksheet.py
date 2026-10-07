"""``AsyncWorksheet`` API: in-memory access, coarse reads, ``run``, the blocking guard."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl.worksheet.worksheet import Worksheet

import aioopenpyxl
from aioopenpyxl import AsyncWorkbook, AsyncWorksheet, BlockingCallError
from tests.conftest import ROWS

pytestmark = pytest.mark.anyio


async def test_fetch_on_read_only_sheet(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        ws = wb["Data"]
        cell = await ws.fetch("B2")
        assert cell.value == "user1"
        block = await ws.fetch("A1:B2")
        assert [[c.value for c in r] for r in block] == [["id", "name"], [1, "user1"]]


async def test_calculate_dimension(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        assert await wb["Data"].calculate_dimension() == f"A1:C{len(ROWS)}"


async def test_calculate_dimension_on_regular_sheet() -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws["B3"] = 1
    assert await ws.calculate_dimension() == "B3:B3"  # openpyxl semantics


async def test_in_memory_cell_access_and_attribute_proxy() -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "hello"
    ws.cell(row=2, column=1, value=42)
    assert ws["A1"].value == "hello"
    assert ws.cell(2, 1).value == 42

    ws.title = "Renamed"
    ws.freeze_panes = "A2"
    ws.merge_cells("B1:C1")
    raw = ws.wrapped
    assert isinstance(raw, Worksheet)
    assert raw.title == "Renamed"
    assert raw.freeze_panes == "A2"
    assert [str(r) for r in ws.merged_cells.ranges] == ["B1:C1"]
    assert ws.max_row == 2
    assert repr(ws) == "<AsyncWorksheet 'Renamed'>"

    # ``__getattr__`` is runtime-only (hidden from type checkers), hence getattr().
    name = "no_such_attribute"
    with pytest.raises(AttributeError):
        getattr(ws, name)


async def test_parent_returns_the_async_workbook() -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.create_sheet("S")
    assert ws.parent is wb
    assert wb["S"].parent is wb
    # A sheet wrapped by hand still gets a working async parent sharing its runner.
    loose = AsyncWorksheet(ws.wrapped)
    parent = loose.parent
    assert isinstance(parent, AsyncWorkbook)
    assert parent.wrapped is wb.wrapped
    assert parent._runner is loose._runner
    assert parent.sheetnames == ["Sheet", "S"]


async def test_append_on_regular_sheet_is_in_memory() -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    await ws.append([1, 2])
    await ws.append({"A": 3, "B": 4})
    assert await ws.read_rows() == [(1, 2), (3, 4)]


async def test_write_only_sheet_close_is_async(tmp_path) -> None:
    wb = aioopenpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("W")
    await ws.append([1])
    assert ws.closed is False
    await ws.close()
    assert ws.closed is True


def _total_of_first_column(raw: Any) -> int:
    # ``raw`` is a ReadOnlyWorksheet at runtime; typed Any because the static union
    # also covers WriteOnlyWorksheet, which has no ``iter_rows``.
    return sum(v for (v,) in raw.iter_rows(min_row=2, max_col=1, values_only=True))


async def test_run_against_raw_worksheet(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        total = await wb["Data"].run(_total_of_first_column)
    assert total == sum(range(1, 51))


def _first_cell_value(raw: Any) -> Any:
    return raw.cell(1, 1).value  # fine off-loop


async def test_cell_on_read_only_sheet_raises_blocking_call_error(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        ws = wb["Data"]
        with pytest.raises(BlockingCallError, match="run"):
            ws.cell(1, 1)
        with pytest.raises(BlockingCallError):
            ws.cell(row=1, column=1)
        # The off-loop alternatives work.
        assert await ws.run(_first_cell_value) == "id"
        assert (await ws.fetch("A1")).value == "id"


async def test_getitem_on_read_only_sheet_raises_blocking_call_error(sample_xlsx: Path) -> None:
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as wb:
        ws = wb["Data"]
        with pytest.raises(BlockingCallError, match="fetch"):
            _ = ws["A1"]
        assert (await ws.fetch("A1")).value == "id"
        raw: Any = ws.wrapped
        assert raw["A1"].value == "id"  # explicit escape hatch still works


async def test_sync_iteration_over_a_worksheet_is_a_type_error(sample_xlsx: Path) -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    await ws.append_rows([[1], [2]])
    with pytest.raises(TypeError, match="async for"):
        for _ in ws:
            pass
    with pytest.raises(TypeError, match="async for"):
        list(ws)
    # ``async for`` is the supported spelling.
    assert [tuple(c.value for c in row) async for row in ws] == [(1,), (2,)]
    async with aioopenpyxl.load_workbook(sample_xlsx, read_only=True) as loaded:
        with pytest.raises(TypeError):
            iter(loaded["Data"])
