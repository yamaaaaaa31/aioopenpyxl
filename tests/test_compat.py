"""Full-surface compatibility with openpyxl: members, sub-module aliases, module constants."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import Any, cast

import openpyxl
import openpyxl.styles
import openpyxl.utils.datetime
import openpyxl.worksheet.table
import pytest
from openpyxl.chart import BarChart
from openpyxl.chartsheet import Chartsheet
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet._write_only import WriteOnlyWorksheet
from openpyxl.worksheet.worksheet import Worksheet

import aioopenpyxl
from aioopenpyxl import AsyncChartsheet, AsyncWorkbook, AsyncWorksheet

pytestmark = pytest.mark.anyio

_NOISE = {
    "__dict__",
    "__weakref__",
    "__module__",
    "__firstlineno__",
    "__static_attributes__",
    "__slots__",
    "__annotations__",
    "__slotnames__",  # cache written onto the class by copy()/pickle
}
# Sync iteration is replaced by ``__aiter__``; Serialisable internals are not sheet API.
_EXEMPT: dict[type, set[str]] = {
    Worksheet: {"__iter__"},
    ReadOnlyWorksheet: {"__iter__"},
    WriteOnlyWorksheet: set(),
    openpyxl.Workbook: set(),
    Chartsheet: {
        "__iter__",
        "__add__",
        "__copy__",
        "__attrs__",
        "__elements__",
        "__namespaced__",
        "__nested__",
        "idx_base",
        "namespace",
    },
}


def _public(obj: Any) -> set[str]:
    names = set()
    for name in dir(obj):
        if name in _NOISE or name in dir(object):
            continue
        if name.startswith("_") and not (name.startswith("__") and name.endswith("__")):
            continue
        names.add(name)
    cls = obj if isinstance(obj, type) else type(obj)
    return names - _EXEMPT[cls]


@pytest.mark.parametrize(
    ("raw", "wrapper"),
    [
        (openpyxl.Workbook, AsyncWorkbook),
        (Worksheet, AsyncWorksheet),
        (ReadOnlyWorksheet, AsyncWorksheet),
        (WriteOnlyWorksheet, AsyncWorksheet),
        (Chartsheet, AsyncChartsheet),
    ],
)
def test_every_public_member_is_defined_on_the_wrapper_class(raw: type, wrapper: type) -> None:
    # Class-level check: the ``__getattr__`` fallback does not count.
    missing = sorted(n for n in _public(raw) if not hasattr(wrapper, n))
    assert missing == []


@pytest.fixture
def raw_instances(sample_xlsx: Path) -> Iterator[list[tuple[Any, Any]]]:
    """(raw openpyxl object, wrapper) pairs covering every runtime flavour."""
    wb = openpyxl.Workbook()
    wo = openpyxl.Workbook(write_only=True)
    ro = openpyxl.load_workbook(sample_xlsx, read_only=True)
    pairs: list[tuple[Any, Any]] = [
        (wb, aioopenpyxl.Workbook.wrap(wb)),
        (wb.active, aioopenpyxl.Workbook.wrap(wb).active),
        (wo.create_sheet("W"), aioopenpyxl.Workbook.wrap(wo)["W"]),
        (ro["Data"], aioopenpyxl.Workbook.wrap(ro)["Data"]),
        (wb.create_chartsheet("C"), aioopenpyxl.Workbook.wrap(wb).chartsheets[0]),
    ]
    try:
        yield pairs
    finally:
        ro.close()


def test_every_public_member_of_live_instances_is_defined_on_the_wrapper_class(
    raw_instances: list[tuple[Any, Any]],
) -> None:
    # Instance-level ``dir()`` also lists attributes only set in ``__init__``;
    # they must be covered by the typed surface (not the ``__getattr__`` fallback).
    for raw, wrapper in raw_instances:
        missing = sorted(n for n in _public(raw) if not hasattr(type(wrapper), n))
        assert missing == [], f"{type(raw).__name__}: {missing}"


async def test_generated_forwarders_keep_openpyxl_semantics() -> None:
    wb = aioopenpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    await ws.append_rows([[1, 2], [3, 4]])
    ws.insert_rows(1)
    assert await ws.read_rows() == [(None, None), (1, 2), (3, 4)]
    ws.delete_rows(1)
    ws.move_range("A1:B2", rows=1, cols=1)
    assert ws["C3"].value == 4
    raw = ws.wrapped
    assert isinstance(raw, Worksheet)
    ws.sheet_state = "hidden"
    assert raw.sheet_state == "hidden"
    header = ws.oddHeader
    assert header is not None and raw.oddHeader is not None
    header.center.text = "hi"
    assert raw.oddHeader.center.text == "hi"
    ws.column_dimensions["A"].width = 30
    assert raw.column_dimensions["A"].width == 30
    assert ws.BREAK_ROW == 1 and ws.PAPERSIZE_A4 == "9"
    assert ws.max_row == 3 and ws.dimensions == "B2:C3"  # openpyxl semantics after move_range
    assert ws.cell(3, 3).value == 4
    assert ws.cell(row=1, column=1, value="x").value == "x"

    wb.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
    assert wb.wrapped.epoch == openpyxl.utils.datetime.CALENDAR_MAC_1904
    assert wb.mime_type == wb.wrapped.mime_type
    assert wb.named_styles == ["Normal"]
    with pytest.deprecated_call():
        assert wb.get_sheet_names() == ["Sheet"]


def test_openpyxl_submodules_are_aliased() -> None:
    # Aliases are installed at import time via sys.modules; the generated ``.pyi``
    # shims make the same imports resolve for type checkers.
    import aioopenpyxl.chart
    import aioopenpyxl.styles
    from aioopenpyxl.styles import Font
    from aioopenpyxl.utils import get_column_letter
    from aioopenpyxl.utils.cell import coordinate_from_string

    assert Font is openpyxl.styles.Font
    assert get_column_letter is openpyxl.utils.get_column_letter
    assert coordinate_from_string is openpyxl.utils.cell.coordinate_from_string
    assert aioopenpyxl.styles is openpyxl.styles
    assert aioopenpyxl.chart.BarChart is BarChart

    import aioopenpyxl.worksheet.table as t

    assert t is openpyxl.worksheet.table


def test_module_level_names_mirror_openpyxl() -> None:
    assert aioopenpyxl.open is aioopenpyxl.load_workbook
    assert aioopenpyxl.LXML == openpyxl.LXML
    assert aioopenpyxl.DEFUSEDXML == openpyxl.DEFUSEDXML
    assert aioopenpyxl.__version__ == "0.1.0"
    assert aioopenpyxl.Workbook is not openpyxl.Workbook
    # ``__getattr__`` is runtime-only (hidden from type checkers), hence getattr().
    name = "no_such_thing"
    with pytest.raises(AttributeError, match="aioopenpyxl"):
        getattr(aioopenpyxl, name)


def test_dunder_all_lists_the_forwarded_constants() -> None:
    for name in ("LXML", "DEBUG", "DEFUSEDXML", "NUMPY", "__author__", "__license__"):
        assert name in aioopenpyxl.__all__
        assert getattr(aioopenpyxl, name) == getattr(openpyxl, name)
    assert "Workbook" in aioopenpyxl.__all__ and "AsyncWorkbook" in aioopenpyxl.__all__
    assert isinstance(aioopenpyxl.LXML, bool)  # readable outside an event loop


def test_custom_doc_props_is_a_typed_property() -> None:
    assert isinstance(aioopenpyxl.Workbook.custom_doc_props, property)
    assert isinstance(aioopenpyxl.AsyncWorkbook.custom_doc_props, property)
    wb = aioopenpyxl.Workbook()
    # (types-openpyxl does not declare it on Workbook; the generator adds it by hand.)
    assert wb.custom_doc_props is cast(Any, wb.wrapped).custom_doc_props
    assert len(wb.custom_doc_props) == 0
