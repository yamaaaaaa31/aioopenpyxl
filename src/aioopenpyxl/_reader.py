"""``load_workbook`` -- awaitable *and* usable as ``async with``."""

from __future__ import annotations

import contextlib
import io
from collections.abc import Awaitable, Coroutine, Generator
from contextlib import AbstractAsyncContextManager
from os import PathLike
from types import TracebackType
from typing import IO, Any

import anyio
import openpyxl
from anyio import CapacityLimiter

from ._executor import Runner, runner_for
from ._workbook import Workbook

__all__ = ["LoadWorkbookContextManager", "load_workbook", "load_workbook_bytes"]


class LoadWorkbookContextManager(
    Awaitable[Workbook],
    AbstractAsyncContextManager[Workbook],
):
    """Return type of :func:`load_workbook`.

    Supports both styles::

        wb = await aioopenpyxl.load_workbook("in.xlsx")
        ...
        await wb.close()

        async with aioopenpyxl.load_workbook("in.xlsx") as wb:
            ...
    """

    __slots__ = ("_coro", "_wb")

    def __init__(self, coro: Coroutine[Any, Any, Workbook]) -> None:
        self._coro = coro
        self._wb: Workbook | None = None

    def __await__(self) -> Generator[Any, None, Workbook]:
        self._wb = yield from self._coro.__await__()
        return self._wb

    async def __aenter__(self) -> Workbook:
        self._wb = await self._coro
        return self._wb

    async def __aexit__(
        self,
        exc_type: type[BaseException] | None,
        exc: BaseException | None,
        tb: TracebackType | None,
    ) -> None:
        if self._wb is not None:
            wb, self._wb = self._wb, None
            # ``Workbook.__aexit__`` completes the close whatever cancellation
            # is in flight (so no read-only zip archive is leaked).
            await wb.__aexit__(exc_type, exc, tb)


async def _open(
    filename: str | PathLike[str] | IO[bytes],
    limiter: CapacityLimiter | None,
    kwargs: dict[str, Any],
) -> Workbook:
    # The load runs under the workbook's future Runner so that a native
    # asyncio cancellation waits for the thread; a workbook that was produced
    # but has no owner any more is closed instead of leaking its archive.
    runner = Runner(limiter)
    loaded: list[openpyxl.Workbook] = []

    def load() -> openpyxl.Workbook:
        wb = openpyxl.load_workbook(filename, **kwargs)
        loaded.append(wb)
        return wb

    try:
        wb = await runner(load)
    except BaseException:
        if loaded:
            with contextlib.suppress(anyio.get_cancelled_exc_class()):
                await runner.run_uncancellable(loaded[0].close)
        raise
    runner_for(wb, runner=runner)
    return Workbook._from_runner(wb, runner)


def load_workbook(
    filename: str | PathLike[str] | IO[bytes],
    read_only: bool = False,
    keep_vba: bool = False,
    data_only: bool = False,
    keep_links: bool = True,
    rich_text: bool = False,
    *,
    limiter: CapacityLimiter | None = None,
) -> LoadWorkbookContextManager:
    """Open an ``.xlsx`` file without blocking the event loop.

    Accepts the same arguments as :func:`openpyxl.load_workbook` plus an
    optional ``limiter`` (an :class:`anyio.CapacityLimiter`).  The result can be
    awaited directly or used with ``async with`` (which closes the workbook on exit).
    """
    kwargs: dict[str, Any] = {
        "read_only": read_only,
        "keep_vba": keep_vba,
        "data_only": data_only,
        "keep_links": keep_links,
        "rich_text": rich_text,
    }
    return LoadWorkbookContextManager(_open(filename, limiter, kwargs))


def load_workbook_bytes(
    data: bytes | bytearray | memoryview,
    read_only: bool = False,
    keep_vba: bool = False,
    data_only: bool = False,
    keep_links: bool = True,
    rich_text: bool = False,
    *,
    limiter: CapacityLimiter | None = None,
) -> LoadWorkbookContextManager:
    """Open an ``.xlsx`` held in memory (an S3 object, an upload body ...).

    Same as :func:`load_workbook` but takes the file content as bytes.
    """
    return load_workbook(
        io.BytesIO(data),
        read_only=read_only,
        keep_vba=keep_vba,
        data_only=data_only,
        keep_links=keep_links,
        rich_text=rich_text,
        limiter=limiter,
    )
