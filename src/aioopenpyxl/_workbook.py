"""Asynchronous proxy around an openpyxl workbook."""

from __future__ import annotations

import io
import weakref
from collections.abc import Callable, Iterator, Sequence
from os import PathLike
from types import TracebackType
from typing import IO, TYPE_CHECKING, Any, Generic, Literal, TypeVar, cast, overload

import openpyxl
from anyio import CapacityLimiter
from openpyxl.chartsheet.chartsheet import Chartsheet
from typing_extensions import Self

from ._errors import WorkbookBusyError
from ._executor import Runner, runner_for
from ._generated import WorkbookProxy
from ._worksheet import AsyncChartsheet, AsyncWorksheet, WS_co

if TYPE_CHECKING:
    from openpyxl.worksheet._write_only import WriteOnlyWorksheet
    from openpyxl.worksheet.worksheet import Worksheet

__all__ = ["AsyncWorkbook", "Workbook"]

R = TypeVar("R")


class Workbook(WorkbookProxy, Generic[WS_co]):
    """Async-aware proxy for :class:`openpyxl.Workbook`.

    ``Workbook()`` creates a new in-memory workbook with the same signature as
    ``openpyxl.Workbook(write_only=..., iso_dates=...)``; construction is
    purely in-memory, so it is a plain call rather than a coroutine.  Wrap an
    existing openpyxl workbook with :meth:`wrap`.

    The type parameter is the raw worksheet type of the sheets this workbook
    hands out: ``Workbook()`` is a ``Workbook[Worksheet]``,
    ``Workbook(write_only=True)`` a ``Workbook[WriteOnlyWorksheet]`` and
    ``load_workbook(..., read_only=True)`` yields a
    ``Workbook[ReadOnlyWorksheet]``.  A plain ``Workbook`` annotation means
    the union of the three and accepts any of them (:meth:`wrap` returns that
    plain type, since a raw workbook does not say which kind it is; spell
    ``Workbook[Worksheet].wrap(raw)`` when you know).

    Every public openpyxl member is available.  Sheet management
    (``create_sheet``, ``remove``, ``wb["Sheet"]`` ...) is in-memory and
    forwarded directly, returning :class:`AsyncWorksheet` /
    :class:`AsyncChartsheet` wrappers that share this workbook's worker
    thread and lock.  :meth:`save`, :meth:`to_bytes`, :meth:`close` and
    :meth:`run` are coroutines executed off the event loop.  Use
    ``async with`` to close the workbook automatically.

    :class:`AsyncWorkbook` is an alias of this class kept for backwards
    compatibility.
    """

    __slots__ = ("_obj", "_runner", "_sheets")

    _obj: openpyxl.Workbook
    # id(raw sheet) -> wrapper, so ``wb["X"] is wb["X"]`` and wrappers are reused.
    _sheets: weakref.WeakValueDictionary[int, AsyncWorksheet[Any] | AsyncChartsheet]

    @overload
    def __init__(
        self: Workbook[Worksheet],
        write_only: Literal[False] = False,
        iso_dates: bool = False,
        *,
        limiter: CapacityLimiter | None = None,
    ) -> None: ...

    @overload
    def __init__(
        self: Workbook[WriteOnlyWorksheet],
        write_only: Literal[True],
        iso_dates: bool = False,
        *,
        limiter: CapacityLimiter | None = None,
    ) -> None: ...

    @overload
    def __init__(
        self: Workbook[Worksheet | WriteOnlyWorksheet],
        write_only: bool,
        iso_dates: bool = False,
        *,
        limiter: CapacityLimiter | None = None,
    ) -> None: ...

    @overload
    def __init__(
        # A subclass of the plain ``Workbook`` (``class MyWorkbook(Workbook)``)
        # is a ``Workbook[RawWorksheet]`` and is constructed through this one.
        self,
        write_only: bool = False,
        iso_dates: bool = False,
        *,
        limiter: CapacityLimiter | None = None,
    ) -> None: ...

    def __init__(
        self,
        write_only: bool = False,
        iso_dates: bool = False,
        *,
        limiter: CapacityLimiter | None = None,
    ) -> None:
        if isinstance(write_only, openpyxl.Workbook):
            # The pre-1.0 ``AsyncWorkbook(raw_workbook)`` call shape; without this
            # guard openpyxl would read the object as ``write_only=True`` and a
            # brand-new empty workbook would be returned silently.
            raise TypeError(
                "Workbook() creates a new workbook; to wrap an existing "
                "openpyxl.Workbook use Workbook.wrap(workbook, limiter=...)"
            )
        raw = openpyxl.Workbook(write_only=write_only, iso_dates=iso_dates)
        self._init(raw, runner_for(raw, limiter))

    @classmethod
    def wrap(
        cls,
        workbook: openpyxl.Workbook,
        *,
        limiter: CapacityLimiter | None = None,
    ) -> Self:
        """Wrap an existing :class:`openpyxl.Workbook` in a fresh :class:`Workbook`.

        A raw workbook has exactly one worker-thread lock (:class:`Runner`),
        shared by every wrapper of it -- wrapping the same workbook twice, or
        creating ``AsyncWorksheet(raw.active)`` by hand, cannot defeat the
        busy guard.  ``limiter`` bounds the threads the workbook may occupy
        (``None`` = anyio's default limiter, or whatever an earlier wrapper
        chose); a *different* limiter than an existing wrapper's raises
        :class:`ValueError`.
        """
        return cls._from_runner(workbook, runner_for(workbook, limiter))

    @classmethod
    def _from_runner(cls, workbook: openpyxl.Workbook, runner: Runner) -> Self:
        self = cls.__new__(cls)
        self._init(workbook, runner_for(workbook, runner=runner))
        return self

    def _init(self, workbook: openpyxl.Workbook, runner: Runner) -> None:
        object.__setattr__(self, "_obj", workbook)
        object.__setattr__(self, "_runner", runner)
        object.__setattr__(self, "_sheets", weakref.WeakValueDictionary())

    # ------------------------------------------------------------------ basics
    @property
    def wrapped(self) -> openpyxl.Workbook:
        """The underlying openpyxl workbook (no busy check)."""
        return self._obj

    @property
    def limiter(self) -> CapacityLimiter | None:
        """The :class:`anyio.CapacityLimiter` bounding this workbook's worker threads."""
        return self._runner.limiter

    def __repr__(self) -> str:
        try:
            return f"<Workbook sheets={self._raw.sheetnames!r}>"
        except WorkbookBusyError:
            return "<Workbook busy>"

    # ------------------------------------------------------- sheet management
    def _wrap(self, sheet: Any) -> AsyncWorksheet[Any] | AsyncChartsheet:
        cached = self._sheets.get(id(sheet))
        if cached is not None and cached.wrapped is sheet:
            return cached
        wrapper: AsyncWorksheet[Any] | AsyncChartsheet
        if isinstance(sheet, Chartsheet):
            wrapper = AsyncChartsheet(sheet, self._runner, self)
        else:
            wrapper = AsyncWorksheet(sheet, self._runner, self)
        self._sheets[id(sheet)] = wrapper
        return wrapper

    @staticmethod
    def _unwrap(sheet: Any) -> Any:
        return sheet.wrapped if isinstance(sheet, AsyncWorksheet | AsyncChartsheet) else sheet

    # Like types-openpyxl, sheets are typed as worksheets for ergonomics even
    # though chartsheets are possible at runtime.
    @property
    def active(self) -> AsyncWorksheet[WS_co] | None:
        ws = self._raw.active
        return None if ws is None else cast("AsyncWorksheet[WS_co]", self._wrap(ws))

    @active.setter
    def active(self, value: AsyncWorksheet[Any] | AsyncChartsheet | int | Any) -> None:
        self._raw.active = self._unwrap(value)

    @property
    def sheetnames(self) -> list[str]:
        return self._raw.sheetnames

    @property
    def worksheets(self) -> Sequence[AsyncWorksheet[WS_co]]:
        # A fresh list each time; typed as ``Sequence`` because ``list`` is
        # invariant and would defeat the covariance of the type parameter.
        return [cast("AsyncWorksheet[WS_co]", self._wrap(ws)) for ws in self._raw.worksheets]

    @property
    def chartsheets(self) -> list[AsyncChartsheet]:
        return [cast(AsyncChartsheet, self._wrap(cs)) for cs in self._raw.chartsheets]

    def __getitem__(self, key: str) -> AsyncWorksheet[WS_co]:
        return cast("AsyncWorksheet[WS_co]", self._wrap(self._raw[key]))

    def __delitem__(self, key: str) -> None:
        del self._raw[key]

    def __contains__(self, key: str) -> bool:
        return key in self._raw

    def __iter__(self) -> Iterator[AsyncWorksheet[WS_co]]:
        return (cast("AsyncWorksheet[WS_co]", self._wrap(ws)) for ws in self._raw)

    def __len__(self) -> int:
        return len(self._raw.worksheets)

    def create_sheet(
        self, title: str | None = None, index: int | None = None
    ) -> AsyncWorksheet[WS_co]:
        return cast(
            "AsyncWorksheet[WS_co]", self._wrap(self._raw.create_sheet(title=title, index=index))
        )

    def create_chartsheet(
        self, title: str | None = None, index: int | None = None
    ) -> AsyncChartsheet:
        return cast(
            AsyncChartsheet, self._wrap(self._raw.create_chartsheet(title=title, index=index))
        )

    def get_sheet_by_name(self, name: str) -> AsyncWorksheet[WS_co]:
        """Deprecated (emits openpyxl's ``DeprecationWarning``); use ``wb[name]``."""
        return cast("AsyncWorksheet[WS_co]", self._wrap(self._raw.get_sheet_by_name(name)))

    def get_sheet_names(self) -> list[str]:
        """Deprecated (emits openpyxl's ``DeprecationWarning``); use ``wb.sheetnames``."""
        return self._raw.get_sheet_names()

    def remove(self, worksheet: AsyncWorksheet[Any] | AsyncChartsheet | Any) -> None:
        self._raw.remove(self._unwrap(worksheet))

    def remove_sheet(self, worksheet: AsyncWorksheet[Any] | AsyncChartsheet | Any) -> None:
        """Deprecated (emits openpyxl's ``DeprecationWarning``); use ``wb.remove(ws)``."""
        self._raw.remove_sheet(self._unwrap(worksheet))

    def index(self, worksheet: AsyncWorksheet[Any] | AsyncChartsheet | Any) -> int:
        return self._raw.index(self._unwrap(worksheet))

    def get_index(self, worksheet: AsyncWorksheet[Any] | AsyncChartsheet | Any) -> int:
        """Deprecated (emits openpyxl's ``DeprecationWarning``); use ``wb.index(ws)``."""
        return self._raw.get_index(self._unwrap(worksheet))

    def copy_worksheet(self, from_worksheet: AsyncWorksheet[Any] | Any) -> AsyncWorksheet[WS_co]:
        return cast(
            "AsyncWorksheet[WS_co]",
            self._wrap(self._raw.copy_worksheet(self._unwrap(from_worksheet))),
        )

    def move_sheet(
        self, sheet: AsyncWorksheet[Any] | AsyncChartsheet | str | Any, offset: int = 0
    ) -> None:
        self._raw.move_sheet(self._unwrap(sheet), offset=offset)

    # ------------------------------------------------------------- async I/O
    async def run(self, func: Callable[[openpyxl.Workbook], R]) -> R:
        """Run ``func(raw_workbook)`` in the worker thread and return its result.

        One thread hop for bulk in-memory work; the per-workbook lock
        guarantees nothing else touches the workbook meanwhile.
        """
        return await self._runner(func, self._obj)

    async def save(self, filename: str | PathLike[str] | IO[bytes]) -> None:
        """Serialise the workbook to ``filename`` (path or binary file object) off the loop."""
        await self._runner(self._obj.save, filename)

    async def to_bytes(self) -> bytes:
        """Serialise the workbook and return the ``.xlsx`` bytes (one thread hop)."""

        def _dump(wb: openpyxl.Workbook) -> bytes:
            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue()

        return await self._runner(_dump, self._obj)

    async def close(self) -> None:
        """Release underlying resources (the zip archive of read-only workbooks).

        Exactly ``openpyxl.Workbook.close``: the in-memory copy a
        ``keep_vba=True`` workbook holds in ``vba_archive`` stays open so a
        later :meth:`save` can still read it.  Leaving the ``async with``
        block closes that copy too (see :meth:`__aexit__`).
        """
        await self._runner(self._obj.close)

    async def __aenter__(self) -> Self:
        return self

    async def __aexit__(
        self,
        exc_type: type[BaseException] | None,
        exc: BaseException | None,
        tb: TracebackType | None,
    ) -> None:
        # Cancellation -- anyio style or native asyncio, even while waiting
        # for the semaphore -- cannot skip the close (which would leak the zip
        # archive of a read-only workbook); it is re-raised afterwards.
        await self._runner.run_uncancellable(_close_for_exit, self._obj)


def _close_for_exit(wb: openpyxl.Workbook) -> None:
    """``wb.close()`` plus the ``vba_archive`` copy of a ``keep_vba=True`` workbook.

    openpyxl keeps that copy as an append-mode ``ZipFile`` over a ``BytesIO``
    and never closes it; when both are reclaimed together (typically at
    interpreter shutdown) the ``ZipFile`` finaliser can run after the buffer
    was closed and print ``ValueError: I/O operation on closed file``.  The
    workbook's lifetime ends with the ``async with`` block, so closing the
    copy here is safe; ``close()`` alone keeps openpyxl's semantics.
    """
    try:
        wb.close()
    finally:
        vba_archive = wb.vba_archive
        if vba_archive is not None:
            vba_archive.close()


#: Backwards-compatible alias; :class:`Workbook` is the canonical name.
AsyncWorkbook = Workbook
