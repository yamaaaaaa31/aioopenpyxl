"""Typed synchronous forwarders for the complete openpyxl surface.

GENERATED FILE -- do not edit.  Regenerate with ``uv run python tools/gen_proxies.py``.
Source: types-openpyxl 3.1.5.20260827 stubs for openpyxl 3.1.5.
The signatures below are derived from the types-openpyxl stubs maintained in typeshed
(https://github.com/python/typeshed, Apache-2.0); openpyxl itself is MIT-licensed and is
used as a dependency, not vendored.  aioopenpyxl is not affiliated with either project.
"""
# ruff: noqa
# fmt: off

from __future__ import annotations

from typing import TYPE_CHECKING, Any, overload

from typing_extensions import deprecated

from ._base import GuardedProxy

if TYPE_CHECKING:
    from _typeshed import ConvertibleToInt, Incomplete, SupportsGetItem, Unused
    from collections.abc import Generator, Iterable, Iterator
    from datetime import datetime
    from openpyxl import _Decodable, _VisibilityType, _ZipFileFileWriteProtocol
    from openpyxl.cell import _AnyCellValue, _CellGetValue, _CellOrMergedCell, _CellSetValue
    from openpyxl.cell.cell import Cell
    from openpyxl.chart._chart import ChartBase
    from openpyxl.chartsheet.chartsheet import Chartsheet
    from openpyxl.chartsheet.custom import CustomChartsheetViews
    from openpyxl.chartsheet.properties import ChartsheetProperties
    from openpyxl.chartsheet.protection import ChartsheetProtection
    from openpyxl.chartsheet.publish import WebPublishItems
    from openpyxl.chartsheet.relation import DrawingHF, SheetBackgroundPicture
    from openpyxl.chartsheet.views import ChartsheetViewList
    from openpyxl.descriptors.base import Alias, Set, Typed
    from openpyxl.descriptors.excel import ExtensionList
    from openpyxl.descriptors.serialisable import Serialisable
    from openpyxl.drawing.image import Image
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.packaging.custom import CustomPropertyList
    from openpyxl.styles.named_styles import NamedStyle
    from openpyxl.utils.cell import _RangeBoundariesTuple
    from openpyxl.utils.indexed_list import IndexedList
    from openpyxl.workbook.child import _WorkbookChild
    from openpyxl.workbook.defined_name import DefinedNameDict
    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    from openpyxl.worksheet._write_only import WriteOnlyWorksheet
    from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
    from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList
    from openpyxl.worksheet.dimensions import ColumnDimension, DimensionHolder, RowDimension, SheetFormatProperties
    from openpyxl.worksheet.drawing import Drawing
    from openpyxl.worksheet.filters import AutoFilter
    from openpyxl.worksheet.header_footer import HeaderFooter, HeaderFooter as _HeaderFooter, HeaderFooterItem
    from openpyxl.worksheet.page import PageMargins, PrintOptions, PrintPageSetup
    from openpyxl.worksheet.pagebreak import ColBreak, RowBreak
    from openpyxl.worksheet.properties import WorksheetProperties
    from openpyxl.worksheet.protection import SheetProtection
    from openpyxl.worksheet.scenario import ScenarioList
    from openpyxl.worksheet.table import Table, TableList
    from openpyxl.worksheet.views import SheetView, SheetViewList
    from openpyxl.worksheet.worksheet import Worksheet
    from openpyxl.xml.functions import Element
    from re import Pattern
    from types import GeneratorType
    from typing import ClassVar, Final, Literal, TypeAlias, type_check_only
    from typing_extensions import Never
    from zipfile import ZipFile


class WorkbookProxy(GuardedProxy):
    """Typed synchronous forwarders generated from types-openpyxl (Workbook)."""

    __slots__ = ()


    @property
    def template(self) -> bool:
        return self._raw.template

    @template.setter
    def template(self, value: bool) -> None:
        self._raw.template = value

    @property
    def path(self) -> str:
        return self._raw.path

    @path.setter
    def path(self, value: str) -> None:
        self._raw.path = value

    @property
    def defined_names(self) -> Incomplete:
        return self._raw.defined_names

    @defined_names.setter
    def defined_names(self, value: Incomplete) -> None:
        self._raw.defined_names = value

    @property
    def properties(self) -> Incomplete:
        return self._raw.properties

    @properties.setter
    def properties(self, value: Incomplete) -> None:
        self._raw.properties = value

    @property
    def security(self) -> Incomplete:
        return self._raw.security

    @security.setter
    def security(self, value: Incomplete) -> None:
        self._raw.security = value

    @property
    def shared_strings(self) -> IndexedList[str]:
        return self._raw.shared_strings

    @shared_strings.setter
    def shared_strings(self, value: IndexedList[str]) -> None:
        self._raw.shared_strings = value

    @property
    def loaded_theme(self) -> Incomplete:
        return self._raw.loaded_theme

    @loaded_theme.setter
    def loaded_theme(self, value: Incomplete) -> None:
        self._raw.loaded_theme = value

    @property
    def vba_archive(self) -> ZipFile | None:
        return self._raw.vba_archive

    @vba_archive.setter
    def vba_archive(self, value: ZipFile | None) -> None:
        self._raw.vba_archive = value

    @property
    def is_template(self) -> bool:
        return self._raw.is_template

    @is_template.setter
    def is_template(self, value: bool) -> None:
        self._raw.is_template = value

    @property
    def code_name(self) -> Incomplete:
        return self._raw.code_name

    @code_name.setter
    def code_name(self, value: Incomplete) -> None:
        self._raw.code_name = value

    @property
    def encoding(self) -> str:
        return self._raw.encoding

    @encoding.setter
    def encoding(self, value: str) -> None:
        self._raw.encoding = value

    @property
    def iso_dates(self) -> Incomplete:
        return self._raw.iso_dates

    @iso_dates.setter
    def iso_dates(self, value: Incomplete) -> None:
        self._raw.iso_dates = value

    @property
    def rels(self) -> Incomplete:
        return self._raw.rels

    @rels.setter
    def rels(self, value: Incomplete) -> None:
        self._raw.rels = value

    @property
    def calculation(self) -> Incomplete:
        return self._raw.calculation

    @calculation.setter
    def calculation(self, value: Incomplete) -> None:
        self._raw.calculation = value

    @property
    def views(self) -> Incomplete:
        return self._raw.views

    @views.setter
    def views(self, value: Incomplete) -> None:
        self._raw.views = value

    @property
    def epoch(self) -> datetime:
        return self._raw.epoch

    @epoch.setter
    def epoch(self, value: datetime) -> None:
        self._raw.epoch = value

    @property
    def read_only(self) -> bool:
        return self._raw.read_only

    @property
    def data_only(self) -> bool:
        return self._raw.data_only

    @property
    def write_only(self) -> bool:
        return self._raw.write_only

    @property
    def excel_base_date(self) -> datetime:
        return self._raw.excel_base_date

    @deprecated('Assign scoped named ranges directly to worksheets or global ones to the workbook. Deprecated in 3.1')
    def create_named_range(self, name: str, worksheet: _WorkbookChild | ReadOnlyWorksheet | None = None, value: str | Incomplete | None = None, scope: Unused = None) -> None:
        return self._raw.create_named_range(name=name, worksheet=worksheet, value=value, scope=scope)

    def add_named_style(self, style: NamedStyle) -> None:
        return self._raw.add_named_style(style=style)

    @property
    def named_styles(self) -> list[str]:
        return self._raw.named_styles

    @property
    def mime_type(self) -> str:
        return self._raw.mime_type

    @property
    def style_names(self) -> list[str]:
        return self._raw.style_names

    @property
    def custom_doc_props(self) -> CustomPropertyList[Any]:
        return self._raw.custom_doc_props

    @custom_doc_props.setter
    def custom_doc_props(self, value: CustomPropertyList[Any]) -> None:
        self._raw.custom_doc_props = value


class WorksheetProxy(GuardedProxy):
    """Typed synchronous forwarders generated from types-openpyxl (Worksheet | ReadOnlyWorksheet | WriteOnlyWorksheet)."""

    __slots__ = ()

    BREAK_NONE: Final = 0
    BREAK_ROW: Final = 1
    BREAK_COLUMN: Final = 2
    SHEETSTATE_VISIBLE: Final = 'visible'
    SHEETSTATE_HIDDEN: Final = 'hidden'
    SHEETSTATE_VERYHIDDEN: Final = 'veryHidden'
    PAPERSIZE_LETTER: Final = '1'
    PAPERSIZE_LETTER_SMALL: Final = '2'
    PAPERSIZE_TABLOID: Final = '3'
    PAPERSIZE_LEDGER: Final = '4'
    PAPERSIZE_LEGAL: Final = '5'
    PAPERSIZE_STATEMENT: Final = '6'
    PAPERSIZE_EXECUTIVE: Final = '7'
    PAPERSIZE_A3: Final = '8'
    PAPERSIZE_A4: Final = '9'
    PAPERSIZE_A4_SMALL: Final = '10'
    PAPERSIZE_A5: Final = '11'
    ORIENTATION_PORTRAIT: Final = 'portrait'
    ORIENTATION_LANDSCAPE: Final = 'landscape'

    @property
    def HeaderFooter(self) -> Any:
        return self._raw.HeaderFooter

    @HeaderFooter.setter
    def HeaderFooter(self, value: Any) -> None:
        self._raw.HeaderFooter = value

    @property
    def encoding(self) -> str:
        return self._raw.encoding

    @property
    def title(self) -> str:
        return self._raw.title

    @title.setter
    def title(self, value: str | _Decodable) -> None:
        self._raw.title = value

    @property
    def oddHeader(self) -> HeaderFooterItem | None:
        return self._raw.oddHeader

    @oddHeader.setter
    def oddHeader(self, value: HeaderFooterItem | None) -> None:
        self._raw.oddHeader = value

    @property
    def oddFooter(self) -> HeaderFooterItem | None:
        return self._raw.oddFooter

    @oddFooter.setter
    def oddFooter(self, value: HeaderFooterItem | None) -> None:
        self._raw.oddFooter = value

    @property
    def evenHeader(self) -> HeaderFooterItem | None:
        return self._raw.evenHeader

    @evenHeader.setter
    def evenHeader(self, value: HeaderFooterItem | None) -> None:
        self._raw.evenHeader = value

    @property
    def evenFooter(self) -> HeaderFooterItem | None:
        return self._raw.evenFooter

    @evenFooter.setter
    def evenFooter(self, value: HeaderFooterItem | None) -> None:
        self._raw.evenFooter = value

    @property
    def firstHeader(self) -> HeaderFooterItem | None:
        return self._raw.firstHeader

    @firstHeader.setter
    def firstHeader(self, value: HeaderFooterItem | None) -> None:
        self._raw.firstHeader = value

    @property
    def firstFooter(self) -> HeaderFooterItem | None:
        return self._raw.firstFooter

    @firstFooter.setter
    def firstFooter(self, value: HeaderFooterItem | None) -> None:
        self._raw.firstFooter = value

    @property
    def path(self) -> str:
        return self._raw.path

    @property
    def mime_type(self) -> str:
        return self._raw.mime_type

    @mime_type.setter
    def mime_type(self, value: str) -> None:
        self._raw.mime_type = value

    @property
    def row_dimensions(self) -> DimensionHolder[int, RowDimension]:
        return self._raw.row_dimensions

    @row_dimensions.setter
    def row_dimensions(self, value: DimensionHolder[int, RowDimension]) -> None:
        self._raw.row_dimensions = value

    @property
    def column_dimensions(self) -> DimensionHolder[str, ColumnDimension]:
        return self._raw.column_dimensions

    @column_dimensions.setter
    def column_dimensions(self, value: DimensionHolder[str, ColumnDimension]) -> None:
        self._raw.column_dimensions = value

    @property
    def row_breaks(self) -> RowBreak:
        return self._raw.row_breaks

    @row_breaks.setter
    def row_breaks(self, value: RowBreak) -> None:
        self._raw.row_breaks = value

    @property
    def col_breaks(self) -> ColBreak:
        return self._raw.col_breaks

    @col_breaks.setter
    def col_breaks(self, value: ColBreak) -> None:
        self._raw.col_breaks = value

    @property
    def merged_cells(self) -> MultiCellRange:
        return self._raw.merged_cells

    @merged_cells.setter
    def merged_cells(self, value: MultiCellRange) -> None:
        self._raw.merged_cells = value

    @property
    def data_validations(self) -> DataValidationList:
        return self._raw.data_validations

    @data_validations.setter
    def data_validations(self, value: DataValidationList) -> None:
        self._raw.data_validations = value

    @property
    def sheet_state(self) -> _VisibilityType:
        return self._raw.sheet_state

    @sheet_state.setter
    def sheet_state(self, value: _VisibilityType) -> None:
        self._raw.sheet_state = value

    @property
    def page_setup(self) -> PrintPageSetup:
        return self._raw.page_setup

    @page_setup.setter
    def page_setup(self, value: PrintPageSetup) -> None:
        self._raw.page_setup = value

    @property
    def print_options(self) -> PrintOptions:
        return self._raw.print_options

    @print_options.setter
    def print_options(self, value: PrintOptions) -> None:
        self._raw.print_options = value

    @property
    def page_margins(self) -> PageMargins:
        return self._raw.page_margins

    @page_margins.setter
    def page_margins(self, value: PageMargins) -> None:
        self._raw.page_margins = value

    @property
    def views(self) -> SheetViewList:
        return self._raw.views

    @views.setter
    def views(self, value: SheetViewList) -> None:
        self._raw.views = value

    @property
    def protection(self) -> SheetProtection:
        return self._raw.protection

    @protection.setter
    def protection(self, value: SheetProtection) -> None:
        self._raw.protection = value

    @property
    def defined_names(self) -> DefinedNameDict:
        return self._raw.defined_names

    @defined_names.setter
    def defined_names(self, value: DefinedNameDict) -> None:
        self._raw.defined_names = value

    @property
    def auto_filter(self) -> AutoFilter:
        return self._raw.auto_filter

    @auto_filter.setter
    def auto_filter(self, value: AutoFilter) -> None:
        self._raw.auto_filter = value

    @property
    def conditional_formatting(self) -> ConditionalFormattingList:
        return self._raw.conditional_formatting

    @conditional_formatting.setter
    def conditional_formatting(self, value: ConditionalFormattingList) -> None:
        self._raw.conditional_formatting = value

    @property
    def legacy_drawing(self) -> Incomplete | None:
        return self._raw.legacy_drawing

    @legacy_drawing.setter
    def legacy_drawing(self, value: Incomplete | None) -> None:
        self._raw.legacy_drawing = value

    @property
    def sheet_properties(self) -> WorksheetProperties:
        return self._raw.sheet_properties

    @sheet_properties.setter
    def sheet_properties(self, value: WorksheetProperties) -> None:
        self._raw.sheet_properties = value

    @property
    def sheet_format(self) -> SheetFormatProperties:
        return self._raw.sheet_format

    @sheet_format.setter
    def sheet_format(self, value: SheetFormatProperties) -> None:
        self._raw.sheet_format = value

    @property
    def scenarios(self) -> ScenarioList:
        return self._raw.scenarios

    @scenarios.setter
    def scenarios(self, value: ScenarioList) -> None:
        self._raw.scenarios = value

    @property
    def sheet_view(self) -> SheetView:
        return self._raw.sheet_view

    @property
    def selected_cell(self) -> str | None:
        return self._raw.selected_cell

    @property
    def active_cell(self) -> str | None:
        return self._raw.active_cell

    @property
    def array_formulae(self) -> dict[str, str]:
        return self._raw.array_formulae

    @property
    def show_gridlines(self) -> bool | None:
        return self._raw.show_gridlines

    @property
    def freeze_panes(self) -> str | None:
        return self._raw.freeze_panes

    @freeze_panes.setter
    def freeze_panes(self, value: str | Cell | None) -> None:
        self._raw.freeze_panes = value

    @property
    def min_row(self) -> int:
        return self._raw.min_row

    @property
    def max_row(self) -> int | None:
        return self._raw.max_row

    @property
    def min_column(self) -> int:
        return self._raw.min_column

    @property
    def max_column(self) -> int | None:
        return self._raw.max_column

    @property
    def dimensions(self) -> str:
        return self._raw.dimensions

    @property
    def column_groups(self) -> list[str]:
        return self._raw.column_groups

    def set_printer_settings(self, paper_size: int | None, orientation: Literal['default', 'portrait', 'landscape'] | None) -> None:
        return self._raw.set_printer_settings(paper_size=paper_size, orientation=orientation)

    def add_data_validation(self, data_validation: DataValidation) -> None:
        return self._raw.add_data_validation(data_validation=data_validation)

    def add_chart(self, chart: ChartBase, anchor: str | None = None) -> None:
        return self._raw.add_chart(chart=chart, anchor=anchor)

    def add_image(self, img: Image, anchor: str | None = None) -> None:
        return self._raw.add_image(img=img, anchor=anchor)

    def add_table(self, table: Table) -> None:
        return self._raw.add_table(table=table)

    @property
    def tables(self) -> TableList:
        return self._raw.tables

    def add_pivot(self, pivot) -> None:
        return self._raw.add_pivot(pivot=pivot)

    @overload
    def merge_cells(self, range_string: str, start_row: None = None, start_column: None = None, end_row: None = None, end_column: None = None) -> None: ...

    @overload
    def merge_cells(self, range_string: None = None, *, start_row: ConvertibleToInt, start_column: ConvertibleToInt, end_row: ConvertibleToInt, end_column: ConvertibleToInt) -> None: ...

    @overload
    def merge_cells(self, range_string: None, start_row: ConvertibleToInt, start_column: ConvertibleToInt, end_row: ConvertibleToInt, end_column: ConvertibleToInt) -> None: ...

    def merge_cells(self, *args: Any, **kwargs: Any) -> Any:
        return self._raw.merge_cells(*args, **kwargs)

    @property
    @deprecated('Use ws.merged_cells.ranges')
    def merged_cell_ranges(self) -> Never:
        self._raw.merged_cell_ranges
        raise AssertionError("openpyxl declares merged_cell_ranges as Never")

    def unmerge_cells(self, range_string: str | None = None, start_row: int | None = None, start_column: int | None = None, end_row: int | None = None, end_column: int | None = None) -> None:
        return self._raw.unmerge_cells(range_string=range_string, start_row=start_row, start_column=start_column, end_row=end_row, end_column=end_column)

    def insert_rows(self, idx: int, amount: int = 1) -> None:
        return self._raw.insert_rows(idx=idx, amount=amount)

    def insert_cols(self, idx: int, amount: int = 1) -> None:
        return self._raw.insert_cols(idx=idx, amount=amount)

    def delete_rows(self, idx: int, amount: int = 1) -> None:
        return self._raw.delete_rows(idx=idx, amount=amount)

    def delete_cols(self, idx: int, amount: int = 1) -> None:
        return self._raw.delete_cols(idx=idx, amount=amount)

    def move_range(self, cell_range: CellRange | str, rows: int = 0, cols: int = 0, translate: bool = False) -> None:
        return self._raw.move_range(cell_range=cell_range, rows=rows, cols=cols, translate=translate)

    @property
    def print_title_rows(self) -> str | None:
        return self._raw.print_title_rows

    @print_title_rows.setter
    def print_title_rows(self, value: str | None) -> None:
        self._raw.print_title_rows = value

    @property
    def print_title_cols(self) -> str | None:
        return self._raw.print_title_cols

    @print_title_cols.setter
    def print_title_cols(self, value: str | None) -> None:
        self._raw.print_title_cols = value

    @property
    def print_titles(self) -> str:
        return self._raw.print_titles

    @property
    def print_area(self) -> str:
        return self._raw.print_area

    @print_area.setter
    def print_area(self, value: str | Iterable[str] | None) -> None:
        self._raw.print_area = value

    def reset_dimensions(self) -> None:
        return self._raw.reset_dimensions()

    @property
    def closed(self) -> bool:
        return self._raw.closed


class ChartsheetProxy(GuardedProxy):
    """Typed synchronous forwarders generated from types-openpyxl (Chartsheet)."""

    __slots__ = ()


    @property
    def HeaderFooter(self) -> Any:
        return self._raw.HeaderFooter

    @HeaderFooter.setter
    def HeaderFooter(self, value: Any) -> None:
        self._raw.HeaderFooter = value

    @property
    def encoding(self) -> str:
        return self._raw.encoding

    @property
    def title(self) -> str:
        return self._raw.title

    @title.setter
    def title(self, value: str | _Decodable) -> None:
        self._raw.title = value

    @property
    def oddHeader(self) -> HeaderFooterItem | None:
        return self._raw.oddHeader

    @oddHeader.setter
    def oddHeader(self, value: HeaderFooterItem | None) -> None:
        self._raw.oddHeader = value

    @property
    def oddFooter(self) -> HeaderFooterItem | None:
        return self._raw.oddFooter

    @oddFooter.setter
    def oddFooter(self, value: HeaderFooterItem | None) -> None:
        self._raw.oddFooter = value

    @property
    def evenHeader(self) -> HeaderFooterItem | None:
        return self._raw.evenHeader

    @evenHeader.setter
    def evenHeader(self, value: HeaderFooterItem | None) -> None:
        self._raw.evenHeader = value

    @property
    def evenFooter(self) -> HeaderFooterItem | None:
        return self._raw.evenFooter

    @evenFooter.setter
    def evenFooter(self, value: HeaderFooterItem | None) -> None:
        self._raw.evenFooter = value

    @property
    def firstHeader(self) -> HeaderFooterItem | None:
        return self._raw.firstHeader

    @firstHeader.setter
    def firstHeader(self, value: HeaderFooterItem | None) -> None:
        self._raw.firstHeader = value

    @property
    def firstFooter(self) -> HeaderFooterItem | None:
        return self._raw.firstFooter

    @firstFooter.setter
    def firstFooter(self, value: HeaderFooterItem | None) -> None:
        self._raw.firstFooter = value

    @property
    def path(self) -> str:
        return self._raw.path

    @property
    def tagname(self) -> str:
        return self._raw.tagname

    @tagname.setter
    def tagname(self, value: str) -> None:
        self._raw.tagname = value

    @property
    def mime_type(self) -> str:
        return self._raw.mime_type

    @mime_type.setter
    def mime_type(self, value: str) -> None:
        self._raw.mime_type = value

    @property
    def sheetPr(self) -> ChartsheetProperties | None:
        return self._raw.sheetPr

    @sheetPr.setter
    def sheetPr(self, value: ChartsheetProperties | None) -> None:
        self._raw.sheetPr = value

    @property
    def sheetViews(self) -> ChartsheetViewList:
        return self._raw.sheetViews

    @sheetViews.setter
    def sheetViews(self, value: ChartsheetViewList) -> None:
        self._raw.sheetViews = value

    @property
    def sheetProtection(self) -> ChartsheetProtection | None:
        return self._raw.sheetProtection

    @sheetProtection.setter
    def sheetProtection(self, value: ChartsheetProtection | None) -> None:
        self._raw.sheetProtection = value

    @property
    def customSheetViews(self) -> CustomChartsheetViews | None:
        return self._raw.customSheetViews

    @customSheetViews.setter
    def customSheetViews(self, value: CustomChartsheetViews | None) -> None:
        self._raw.customSheetViews = value

    @property
    def pageMargins(self) -> PageMargins | None:
        return self._raw.pageMargins

    @pageMargins.setter
    def pageMargins(self, value: PageMargins | None) -> None:
        self._raw.pageMargins = value

    @property
    def pageSetup(self) -> PrintPageSetup | None:
        return self._raw.pageSetup

    @pageSetup.setter
    def pageSetup(self, value: PrintPageSetup | None) -> None:
        self._raw.pageSetup = value

    @property
    def drawing(self) -> Drawing | None:
        return self._raw.drawing

    @drawing.setter
    def drawing(self, value: Drawing | None) -> None:
        self._raw.drawing = value

    @property
    def drawingHF(self) -> DrawingHF | None:
        return self._raw.drawingHF

    @drawingHF.setter
    def drawingHF(self, value: DrawingHF | None) -> None:
        self._raw.drawingHF = value

    @property
    def picture(self) -> SheetBackgroundPicture | None:
        return self._raw.picture

    @picture.setter
    def picture(self, value: SheetBackgroundPicture | None) -> None:
        self._raw.picture = value

    @property
    def webPublishItems(self) -> WebPublishItems | None:
        return self._raw.webPublishItems

    @webPublishItems.setter
    def webPublishItems(self, value: WebPublishItems | None) -> None:
        self._raw.webPublishItems = value

    @property
    def extLst(self) -> ExtensionList | None:
        return self._raw.extLst

    @extLst.setter
    def extLst(self, value: ExtensionList | None) -> None:
        self._raw.extLst = value

    @property
    def sheet_state(self) -> _VisibilityType:
        return self._raw.sheet_state

    @sheet_state.setter
    def sheet_state(self, value: _VisibilityType) -> None:
        self._raw.sheet_state = value

    @property
    def headerFooter(self) -> _HeaderFooter:
        return self._raw.headerFooter

    @headerFooter.setter
    def headerFooter(self, value: _HeaderFooter) -> None:
        self._raw.headerFooter = value

    def add_chart(self, chart) -> None:
        return self._raw.add_chart(chart=chart)

    def to_tree(self) -> Element:
        return self._raw.to_tree()

    @property
    def idx_base(self) -> int:
        return self._raw.idx_base

    @idx_base.setter
    def idx_base(self, value: int) -> None:
        self._raw.idx_base = value

    @property
    def namespace(self) -> str | None:
        return self._raw.namespace

    @namespace.setter
    def namespace(self, value: str | None) -> None:
        self._raw.namespace = value
