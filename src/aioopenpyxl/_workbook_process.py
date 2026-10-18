"""One persistent workbook per spawned, single-worker process pool."""

from __future__ import annotations

import contextlib
import inspect
import io
import multiprocessing
import os
import traceback
from collections.abc import Callable
from concurrent.futures import ProcessPoolExecutor
from concurrent.futures.process import BrokenProcessPool
from typing import IO, Any, TypeVar, cast

import cloudpickle  # type: ignore[import-untyped]  # No upstream stubs or py.typed marker.
import openpyxl
from anyio import CapacityLimiter

from ._executor import Runner

R = TypeVar("R")
Spec = tuple[str, str | bytes | None, dict[str, Any]]

# Each instance owns a single-worker executor. Only the child accesses this
# global, so subsequent commands reach the same live workbook/open archive.
_workbook: openpyxl.Workbook | None = None


def _close_workbook(wb: openpyxl.Workbook) -> None:
    from ._workbook import _close_for_exit

    try:
        if wb.write_only:
            for ws in wb.worksheets:
                writer = getattr(ws, "_writer", None)
                if writer is not None:
                    try:
                        if not ws.closed:
                            ws.close()
                    finally:
                        if os.path.exists(writer.out):
                            writer.cleanup()
    finally:
        _close_for_exit(wb)


def _dispatch(spec: Spec, operation: str, payload: Any = None) -> bytes:
    global _workbook
    if operation == "close":
        if _workbook is not None:
            try:
                _close_workbook(_workbook)
            finally:
                _workbook = None
        return b""
    if _workbook is None:
        kind, source, options = spec
        if kind == "new":
            _workbook = openpyxl.Workbook(**options)
        else:
            assert source is not None
            _workbook = openpyxl.load_workbook(
                io.BytesIO(source) if isinstance(source, bytes) else source, **options
            )
    if operation == "open":
        return b""
    if operation == "run":
        try:
            func = cloudpickle.loads(payload)
            result = func(_workbook)
            if inspect.isawaitable(result):
                if inspect.iscoroutine(result):
                    result.close()
                raise TypeError("Workbook.run() requires a synchronous callback")
            return cloudpickle.dumps((True, result, None))
        except BaseException as exc:
            # Preserve locally defined exception classes as well as lambdas.
            return cloudpickle.dumps((False, exc, traceback.format_exc()))
    if operation == "save":
        _workbook.save(payload)
        return b""
    if operation == "bytes":
        output = io.BytesIO()
        _workbook.save(output)
        return output.getvalue()
    raise ValueError(f"Unknown workbook process operation: {operation}")


class WorkbookProcess:
    """Private process backend; every command and close share the same Runner.

    The coordinator thread performs serialization, process startup, submission,
    result decoding and file-like I/O. The child keeps the actual workbook.
    Nothing starts until an awaited operation, so Workbook() remains synchronous.
    """

    def __init__(
        self,
        options: dict[str, Any],
        *,
        source: str | os.PathLike[str] | IO[bytes] | None = None,
        loading: bool = False,
        limiter: CapacityLimiter | None = None,
    ) -> None:
        self.runner = Runner(limiter)
        self._options = options
        self._source = source
        self._loading = loading
        self._spec: Spec | None = None
        self._executor: ProcessPoolExecutor | None = None
        self._closing = False
        self._closed = False

    def _send(self, operation: str, payload: Any = None) -> bytes:
        if self._closing:
            raise RuntimeError("Process workbook is closed")
        if self._spec is None:
            source: str | bytes | None = None
            if self._loading:
                if isinstance(self._source, str | os.PathLike):
                    source = os.path.abspath(self._source)
                else:
                    assert self._source is not None
                    # ZipFile ignores the initial position. Snapshot the whole
                    # seekable stream, then leave its position unchanged.
                    position = self._source.tell()
                    try:
                        self._source.seek(0)
                        source = self._source.read()
                    finally:
                        self._source.seek(position)
                self._source = None
            self._spec = ("load" if self._loading else "new", source, self._options)
        if self._executor is None:
            self._executor = ProcessPoolExecutor(
                max_workers=1, mp_context=multiprocessing.get_context("spawn")
            )
        # Only initialization needs the source data. Later calls must not
        # resend a potentially large upload just to identify a live workbook.
        response = self._executor.submit(_dispatch, self._spec, operation, payload).result()
        self._spec = (self._spec[0], None, self._spec[2])
        return response

    async def open(self) -> None:
        await self.runner(self._send, "open")

    async def run(self, func: Callable[[openpyxl.Workbook], R]) -> R:
        def invoke() -> R:
            callback = cloudpickle.dumps(func)
            success, result, remote_traceback = cloudpickle.loads(self._send("run", callback))
            if not success:
                raise result from RuntimeError(
                    f"Exception in workbook process:\n{remote_traceback}"
                )
            return cast(R, result)

        return await self.runner(invoke)

    async def save(self, filename: str | os.PathLike[str] | IO[bytes]) -> None:
        def save() -> None:
            if isinstance(filename, str | os.PathLike):
                # Callbacks can change the child's cwd, and the parent may
                # change its own cwd after the child started. Match the
                # caller's path semantics rather than the worker's cwd.
                self._send("save", os.path.abspath(filename))
            else:
                filename.write(self._send("bytes"))

        await self.runner(save)

    async def to_bytes(self) -> bytes:
        return await self.runner(self._send, "bytes")

    def _close(self) -> None:
        if self._closed:
            return
        executor = self._executor
        try:
            if executor is not None:
                try:
                    # A crashed worker is already gone; still join the pool,
                    # and do not mask the original BrokenProcessPool error.
                    with contextlib.suppress(BrokenProcessPool):
                        assert self._spec is not None
                        executor.submit(_dispatch, self._spec, "close").result()
                finally:
                    executor.shutdown(wait=True)
        finally:
            self._executor = None
            self._source = None
            self._spec = None
            self._closed = True

    async def close(self) -> None:
        self._closing = True
        await self.runner.run_uncancellable(self._close)
