"""aioopenpyxl -- asynchronous wrapper around openpyxl.

    import aioopenpyxl

    async with aioopenpyxl.load_workbook("data.xlsx", read_only=True) as wb:
        ws = wb.active
        async for row in ws.iter_rows(values_only=True):
            ...

    wb = aioopenpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "header"              # in-memory: forwarded directly
    await ws.append_rows(rows)       # one thread hop for the whole batch
    await wb.save("output.xlsx")     # off the event loop

Built on anyio: works on asyncio and trio alike.

:class:`Workbook` is a class (``AsyncWorkbook`` is a backwards-compatible
alias), so it works in annotations and ``isinstance`` checks and has the same
constructor signature as ``openpyxl.Workbook``; :meth:`Workbook.wrap` wraps an
existing openpyxl workbook.  Every public openpyxl member is available on the
wrappers.  Blocking calls (loading, saving, parsing rows of read-only sheets,
spooling rows of write-only sheets, ``run(...)``) execute in a worker thread
and are serialised per workbook; while one is in flight, synchronous access
raises :class:`WorkbookBusyError` instead of racing.

``import aioopenpyxl as openpyxl`` is supported as a migration path: the
top-level ``Workbook`` / ``load_workbook`` / ``open`` are the *async* wrappers
(``save``/``close``/``load`` must be awaited), while openpyxl's sub-packages
(``aioopenpyxl.styles``, ``aioopenpyxl.utils`` ...) and its module-level
constants (``LXML``, ``DEBUG`` ...) are the very same objects as openpyxl's.
Note that ``aioopenpyxl.workbook.Workbook`` and
``aioopenpyxl.reader.excel.load_workbook`` are therefore the *synchronous*
openpyxl originals; only the top-level names are asynchronous.
"""

from __future__ import annotations

import sys
from typing import TYPE_CHECKING, Any

# Module-level constants openpyxl exposes; re-exported explicitly so that
# static analysis sees them (the ``__getattr__`` fallback below is runtime only).
from openpyxl import (
    DEBUG,
    DEFUSEDXML,
    LXML,
    NUMPY,
    __author__,
    __author_email__,
    __license__,
    __maintainer_email__,
    __url__,
)

from ._compat import install_module_aliases
from ._errors import BlockingCallError, WorkbookBusyError
from ._executor import Runner, run_sync
from ._reader import LoadWorkbookContextManager, load_workbook, load_workbook_bytes
from ._workbook import AsyncWorkbook, Workbook
from ._worksheet import DEFAULT_CHUNK_SIZE, AsyncChartsheet, AsyncWorksheet

#: aioopenpyxl's own version (not openpyxl's).
__version__ = "0.1.1"

#: Alias mirroring ``openpyxl.open``.
open = load_workbook

__all__ = [
    "DEBUG",
    "DEFAULT_CHUNK_SIZE",
    "DEFUSEDXML",
    "LXML",
    "NUMPY",
    "AsyncChartsheet",
    "AsyncWorkbook",
    "AsyncWorksheet",
    "BlockingCallError",
    "LoadWorkbookContextManager",
    "Runner",
    "Workbook",
    "WorkbookBusyError",
    "__author__",
    "__author_email__",
    "__license__",
    "__maintainer_email__",
    "__url__",
    "__version__",
    "load_workbook",
    "load_workbook_bytes",
    "open",
    "run_sync",
]

install_module_aliases(sys.modules[__name__])


if not TYPE_CHECKING:
    # Runtime-only safety net for anything openpyxl may add at module level in
    # a future release.  Hidden from type checkers so that a typo on
    # ``aioopenpyxl.<name>`` is reported instead of silently becoming ``Any``.
    def __getattr__(name: str) -> Any:
        import openpyxl as _openpyxl

        try:
            return getattr(_openpyxl, name)
        except AttributeError:
            raise AttributeError(f"module {__name__!r} has no attribute {name!r}") from None
